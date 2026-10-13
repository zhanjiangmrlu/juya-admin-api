import hashlib
import json
import re
import stat
import tempfile
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Iterator, Mapping
from copy import deepcopy
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from juya_admin_api.modules.content.v15_contracts import IMPORT_NAMESPACE, ContentType
from juya_admin_api.modules.content_import.sources import assign_sentence_ids, resolve_entry_source
from juya_admin_api.shared.errors import AppError


@dataclass(frozen=True)
class ImportLimits:
    max_package_bytes: int = 512 * 1024 * 1024
    max_total_bytes: int = 2 * 1024 * 1024 * 1024
    max_file_bytes: int = 50 * 1024 * 1024
    max_files: int = 12000
    max_ratio: int = 200
    max_path_length: int = 240
    max_items: int = 1000


def canonical_hash(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def semantic_input(value: Any) -> Any:
    """Exclude evidence locations from replay hashes, preserving declared content identity."""
    if isinstance(value, dict):
        return {
            key: semantic_input(item)
            for key, item in value.items()
            if key
            not in {
                "source_snapshot",
                "parsing_issues",
                "parsed_entries",
                "source_locations",
                "input_origin",
                "chinese_speaker",
            }
        }
    if isinstance(value, list):
        return [semantic_input(item) for item in value]
    return value


def safe_path(value: str, limits: ImportLimits) -> str:
    value = unicodedata.normalize("NFC", value.replace("\\", "/"))
    parts = value.split("/")
    if (
        not value
        or len(value) > limits.max_path_length
        or value.startswith("/")
        or any(part in {"", ".", ".."} or part.endswith((" ", ".")) for part in parts)
        or any(c in value for c in (":", "\x00"))
        or any(ord(c) < 32 for c in value)
        or any(
            part.split(".")[0].upper()
            in {
                "CON",
                "PRN",
                "AUX",
                "NUL",
                *(f"COM{i}" for i in range(1, 10)),
                *(f"LPT{i}" for i in range(1, 10)),
            }
            for part in parts
        )
    ):
        raise AppError("IMPORT_PATH_INVALID", "包内路径不安全,请使用明确相对路径", 422)
    return value


class SpoolFiles(Mapping[str, bytes]):
    """Keep expanded batch resources on private disk, reading at most one bounded file at a time."""

    def __init__(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="juya-import-")
        self.paths: dict[str, Path] = {}

    def __getitem__(self, key: str) -> bytes:
        return self.paths[key].read_bytes()

    def __iter__(self) -> Iterator[str]:
        return iter(self.paths)

    def __len__(self) -> int:
        return len(self.paths)

    def close(self) -> None:
        self.temporary.cleanup()


def read_container(
    data: bytes, limits: ImportLimits, *, spool: bool = False
) -> Mapping[str, bytes]:
    if len(data) > limits.max_package_bytes:
        raise AppError("IMPORT_LIMIT_EXCEEDED", "压缩包超过上传限制", 422)
    expanded = SpoolFiles() if spool else None
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_files:
                raise AppError("IMPORT_LIMIT_EXCEEDED", "压缩包文件数量超过限制", 422)
            paths: dict[str, str] = {}
            for info in infos:
                path = safe_path(
                    info.filename.rstrip("/") if info.is_dir() else info.filename, limits
                )
                if (
                    path.casefold() in paths
                    or stat.S_ISLNK(info.external_attr >> 16)
                    or info.flag_bits & 1
                ):
                    raise AppError("IMPORT_PATH_INVALID", "压缩包包含重复路径、链接或加密文件", 422)
                mode = stat.S_IFMT(info.external_attr >> 16)
                if mode not in {0, stat.S_IFREG, stat.S_IFDIR}:
                    raise AppError("IMPORT_PATH_INVALID", "压缩包不能包含设备或特殊文件", 422)
                paths[path.casefold()] = path
                if (
                    info.file_size > limits.max_file_bytes
                    or info.file_size > max(1, info.compress_size) * limits.max_ratio
                ):
                    raise AppError("IMPORT_LIMIT_EXCEEDED", "文件展开规模或压缩比超过限制", 422)
            files: dict[str, bytes] = {}
            total = 0
            for info in infos:
                if info.is_dir():
                    continue
                chunks: list[bytes] = []
                count = 0
                destination = (
                    Path(expanded.temporary.name) / str(len(expanded.paths))
                    if expanded is not None
                    else None
                )
                with archive.open(info) as stream:
                    output = destination.open("wb") if destination is not None else None
                    try:
                        while block := stream.read(65536):
                            count += len(block)
                            total += len(block)
                            if (
                                count > limits.max_file_bytes
                                or total > limits.max_total_bytes
                                or count > max(1, info.compress_size) * limits.max_ratio
                            ):
                                raise AppError("IMPORT_LIMIT_EXCEEDED", "实际展开字节超过限制", 422)
                            if output is not None:
                                output.write(block)
                            else:
                                chunks.append(block)
                    finally:
                        if output is not None:
                            output.close()
                path = safe_path(info.filename, limits)
                if expanded is not None and destination is not None:
                    expanded.paths[path] = destination
                else:
                    files[path] = b"".join(chunks)
            return expanded if expanded is not None else files
    except (BadZipFile, RuntimeError, OSError, EOFError) as error:
        if expanded is not None:
            expanded.close()
        raise AppError("IMPORT_CONTAINER_INVALID", "压缩文件损坏,无法解析", 422) from error
    except BaseException:
        if expanded is not None:
            expanded.close()
        raise


def validate_ooxml(data: bytes, limits: ImportLimits) -> None:
    files = read_container(data, limits)
    for path, body in files.items():
        if "vbaproject" in path.casefold() or path.endswith((".exe", ".dll")):
            raise AppError("IMPORT_DOCUMENT_UNSAFE", "文档不能包含宏或可执行附件", 422)
        if path.endswith((".xml", ".rels")):
            try:
                tree = ElementTree.fromstring(body)
                if path.endswith(".rels") and any(
                    e.get("TargetMode", "").casefold() == "external" for e in tree
                ):
                    raise AppError("IMPORT_DOCUMENT_UNSAFE", "文档不能包含外部关系", 422)
            except AppError:
                raise
            except Exception as error:
                raise AppError(
                    "IMPORT_DOCUMENT_UNSAFE", "文档XML包含不安全或损坏内容", 422
                ) from error


def read_sheet(
    data: bytes, limits: ImportLimits, *, include_empty: bool = False
) -> list[dict[str, Any]]:
    validate_ooxml(data, limits)
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            rows = iter(book.worksheets[0].iter_rows())
            headings = [str(c.value or "").strip() for c in next(rows)]
            if len(headings) != len(set(headings)):
                raise ValueError("duplicate columns")
            result: list[dict[str, Any]] = []
            for number, cells in enumerate(rows, start=2):
                if number > limits.max_items * 100 + 2:
                    raise AppError("IMPORT_LIMIT_EXCEEDED", "表格行数超过限制", 422)
                row: dict[str, Any] = {
                    head: str(cell.value).strip() if cell.value is not None else ""
                    for head, cell in zip(headings, cells, strict=False)
                    if head
                }
                if not include_empty and not any(row.values()):
                    continue
                row["original_row"] = {
                    "row_no": number,
                    "values": {
                        head: cell.value
                        if cell.value is None or isinstance(cell.value, str | int | float | bool)
                        else str(cell.value)
                        for head, cell in zip(headings, cells, strict=False)
                        if head
                    },
                    "cells": [
                        {
                            "column": index,
                            "heading": head,
                            "value": cell.value
                            if cell.value is None
                            or isinstance(cell.value, str | int | float | bool)
                            else str(cell.value),
                            "data_type": cell.data_type,
                        }
                        for index, (head, cell) in enumerate(
                            zip(headings, cells, strict=False), start=1
                        )
                    ],
                }
                row["row_no"] = number
                row["formula_columns"] = [
                    head
                    for head, cell in zip(headings, cells, strict=False)
                    if cell.data_type == "f"
                ]
                result.append(row)
            return result
        finally:
            book.close()
    except AppError:
        raise
    except Exception as error:
        raise AppError("IMPORT_XLSX_INVALID", "内容表无法解析或列名重复", 422) from error


def parsing_issue(code: str, field: str, message: str, locations: list[str]) -> dict[str, Any]:
    return {
        "id": "parse-" + canonical_hash([code, field, locations])[:24],
        "code": code,
        "field": field,
        "message": message,
        "severity": "DRAFT",
        "source_locations": locations,
    }


def speaker_text(text: str) -> tuple[str, str]:
    match = re.match(r"^([^:\uff1a]{1,100})[:\uff1a]\s*(.+)$", text.strip())
    return (match[1].strip(), match[2]) if match else ("", text.strip())


def pair_dialogue(article: dict[str, Any]) -> list[dict[str, Any]]:
    """Keep ordinal grouped candidates; use local adjacency for mixed bilingual input."""
    english, chinese = article["english"], article["chinese"]
    paired: dict[str, dict[str, Any]] = {}
    body = [b for b in article["blocks"] if b["role"] in {"english", "chinese"}]
    languages = [b["role"] for b in body]
    en_names = {speaker_text(en["text"])[0] for en in english} - {""}
    zh_names = {speaker_text(zh["text"])[0] for zh in chinese} - {""}
    shared_names = en_names & zh_names
    grouped = languages == ["english"] * len(english) + ["chinese"] * len(chinese) and (
        article["explicit_languages"]
        or (len(english) == len(chinese) and (not shared_names or en_names == zh_names))
    )
    if grouped:
        paired = {en["id"]: zh for en, zh in zip(english, chinese, strict=False)}
    else:
        previous: dict[str, Any] | None = None
        for block in body:
            if block["role"] == "english":
                previous = block
            elif previous is not None:
                paired[previous["id"]] = block
                previous = None
    consumed = {zh["id"] for zh in paired.values()}
    dialogue = []
    aliases: dict[str, str] = {}
    reverse_aliases: dict[str, str] = {}
    for en in english:
        zh = paired.get(en["id"])
        name, text = speaker_text(en["text"])
        zh_name, translation = speaker_text(zh["text"]) if zh else ("", "")
        locations = [en["id"], *([zh["id"]] if zh else [])]
        if zh is None:
            article["issues"].append(
                parsing_issue(
                    "BILINGUAL_PAIRING_REQUIRED",
                    "dialogue",
                    "此英文缺少可确定对应的中文",
                    locations,
                )
            )
        elif name and zh_name:
            conflict = (
                ((name in shared_names or zh_name in shared_names) and name != zh_name)
                or (name in aliases and aliases[name] != zh_name)
                or (zh_name in reverse_aliases and reverse_aliases[zh_name] != name)
            )
            if conflict:
                article["issues"].append(
                    parsing_issue(
                        "BILINGUAL_SPEAKER_CONFLICT",
                        "dialogue",
                        "说话者名称或重复出现顺序不一致,需校对",
                        locations,
                    )
                )
            aliases.setdefault(name, zh_name)
            reverse_aliases.setdefault(zh_name, name)
        dialogue.append(
            {
                "speaker": name,
                "english": text,
                "chinese": translation,
                "chinese_speaker": zh_name,
                "source_locations": locations,
            }
        )
    for zh in chinese:
        if zh["id"] not in consumed:
            article["issues"].append(
                parsing_issue(
                    "BILINGUAL_EXTRA_CHINESE", "dialogue", "此中文没有可确定对应的英文", [zh["id"]]
                )
            )
    return assign_sentence_ids(dialogue)


LEXICON_COLUMNS = {
    "import_entry_id": "import_entry_id",
    "entry_id": "import_entry_id",
    "条目id": "import_entry_id",
    "english": "english",
    "英文": "english",
    "word": "english",
    "单词": "english",
    "chunk": "english",
    "语块": "english",
    "meaning_zh": "meaning_zh",
    "chinese": "meaning_zh",
    "中文释义": "meaning_zh",
    "释义": "meaning_zh",
    "中文": "meaning_zh",
    "explanation": "explanation",
    "explanation_zh": "explanation",
    "解释": "explanation",
    "phonetic": "phonetic",
    "音标": "phonetic",
    "source_sentence": "source_sentence",
    "来源句": "source_sentence",
    "source_locator": "source_locator",
    "来源位置": "source_locator",
    "entry_type": "entry_type",
    "image_path": "image_path",
    "pronunciation_path": "pronunciation_path",
    "pronunciation_text": "pronunciation_text",
    "reading_variant": "reading_variant",
}


def materialize_docx_entries(article: dict[str, Any], kind: ContentType) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    if kind == "short_reading":
        return entries
    for attachment in article["attachments"]:
        supplied = attachment.get("fields")
        if supplied is None and attachment.get("kind") == "paragraph":
            parts = [p.strip() for p in re.split(r"[|\uff5c\t]", attachment["text"])]
            if 2 <= len(parts) <= 4:
                supplied = dict(
                    zip(
                        ("english", "meaning_zh", "explanation", "source_sentence"),
                        parts,
                        strict=False,
                    )
                )
        if not supplied or not supplied.get("english") or not supplied.get("meaning_zh"):
            article["issues"].append(
                parsing_issue(
                    "IMPORT_LEXICON_AMBIGUOUS",
                    "entries",
                    "词表格式或释义不完整,保留原资料供校对",
                    attachment["source_locations"],
                )
            )
            continue
        entry_type = attachment["entry_type"]
        if supplied.get("entry_type") not in {None, "", entry_type}:
            article["issues"].append(
                parsing_issue(
                    "ENTRY_TYPE_INVALID",
                    "entry_type",
                    "词表行类型与所在章节不一致,需明确确认",
                    attachment["source_locations"],
                )
            )
            continue
        entry = dict(supplied)
        entry.update(
            entry_type=entry_type,
            import_entry_id=supplied.get("import_entry_id")
            or f"docx-{entry_type}-" + canonical_hash([entry_type, supplied])[:32],
            source_locations=attachment["source_locations"],
            input_origin="docx",
        )
        if kind == "dialogue" and (entry.get("source_sentence") or entry.get("source_locator")):
            resolved = resolve_entry_source(entry, article["dialogue"])
            entry.update(
                {
                    k: resolved[k]
                    for k in ("source_sentence", "source_sentence_ids", "source_locator")
                }
            )
            for issue in resolved["issues"]:
                article["issues"].append(
                    issue | {"source_locations": attachment["source_locations"]}
                )
        if entry_type == "chunk" and entry.get("image_path"):
            article["issues"].append(
                parsing_issue(
                    "CHUNK_IMAGE_INVALID",
                    "image_path",
                    "语块不接收图片",
                    attachment["source_locations"],
                )
            )
        entries.append(entry)
    duplicate_keys = Counter(e["import_entry_id"] for e in entries)
    for entry in entries:
        if duplicate_keys[entry["import_entry_id"]] > 1:
            article["issues"].append(
                parsing_issue(
                    "IMPORT_DUPLICATE_ENTRY_ID",
                    "import_entry_id",
                    "DOCX条目ID重复,请明确独立导入ID",
                    entry["source_locations"],
                )
            )
    return [entry for entry in entries if duplicate_keys[entry["import_entry_id"]] == 1]


def original_inner_blocks(parts: Iterator[Paragraph | Table], prefix: str) -> list[dict[str, Any]]:
    """Keep unsupported nested tables and each cell paragraph as addressable original input."""
    result: list[dict[str, Any]] = []
    for index, part in enumerate(parts, start=1):
        location = f"{prefix}-block-{index:04d}"
        if isinstance(part, Paragraph):
            result.append(
                {
                    "id": location,
                    "kind": "paragraph",
                    "text": part.text,
                    "role": "unconsumed",
                }
            )
        else:
            rows = []
            for row_no, row in enumerate(part.rows, start=1):
                cells = []
                for column_no, cell in enumerate(row.cells, start=1):
                    cell_id = f"{location}-row-{row_no:04d}-cell-{column_no:04d}"
                    cells.append(
                        {
                            "id": cell_id,
                            "kind": "table_cell",
                            "text": cell.text,
                            "role": "unconsumed",
                            "children": original_inner_blocks(cell.iter_inner_content(), cell_id),
                        }
                    )
                rows.append(cells)
            result.append({"id": location, "kind": "table", "rows": rows, "role": "unconsumed"})
    return result


def parse_docx(data: bytes, kind: ContentType, limits: ImportLimits) -> list[dict[str, Any]]:
    validate_ooxml(data, limits)
    try:
        doc = Document(BytesIO(data))
    except Exception as error:
        raise AppError("IMPORT_DOCX_INVALID", "双语DOCX无法解析", 422) from error
    articles: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    preamble: list[dict[str, Any]] = []
    mode = "body"
    paragraph_no = table_no = 0

    def new_article() -> dict[str, Any]:
        result: dict[str, Any] = {
            "title_en": "",
            "title_zh": "",
            "source": "",
            "copyright": "",
            "english": [],
            "chinese": [],
            "blocks": list(preamble),
            "attachments": [],
            "issues": [],
            "explicit_languages": False,
        }
        preamble.clear()
        articles.append(result)
        return result

    for part in doc.iter_inner_content():
        if isinstance(part, Paragraph):
            paragraph_no += 1
            text = part.text.strip().replace("\r\n", "\n")
            block: dict[str, Any] = {
                "id": f"paragraph-{paragraph_no:04d}",
                "kind": "paragraph",
                "paragraph_index": paragraph_no,
                "text": part.text,
                "role": "unconsumed",
            }
            if (part.style.name if part.style else "").startswith(
                ("Heading", "标题")
            ) and re.search(r"[\uff5c|]", text):
                current = new_article()
                titles = re.split(r"[\uff5c|]", text, maxsplit=1)
                current.update(title_zh=titles[0].strip(), title_en=titles[1].strip())
                mode = "body"
                block["role"] = "title"
                current["blocks"].append(block)
                continue
            if current is None:
                if not text:
                    preamble.append(block)
                    continue
                current = new_article()
            current["blocks"].append(block)
            if not text:
                block["role"] = "blank"
                continue
            headings = {
                "英文正文": "english",
                "english": "english",
                "english text": "english",
                "中文译文": "chinese",
                "全文中文译文": "chinese",
                "中文正文": "chinese",
                "translation": "chinese",
                "chinese": "chinese",
                "vocabulary": "word",
                "重点词汇": "word",
                "词汇": "word",
                "useful chunks": "chunk",
                "语块": "chunk",
            }
            if text.casefold().rstrip(":\uff1a") in headings:
                mode = headings[text.casefold().rstrip(":\uff1a")]
                block["role"] = "section_heading"
                current["explicit_languages"] |= mode in {"english", "chinese"}
                continue
            metadata = re.match(
                r"^(Source|来源|Copyright|版权|Safe Intro|安全简介|Topic|题材|Level|难度)"
                r"[:\uff1a]\s*(.*)$",
                text,
                re.IGNORECASE,
            )
            if metadata:
                field = {
                    "source": "source",
                    "来源": "source",
                    "copyright": "copyright",
                    "版权": "copyright",
                    "safe intro": "safe_intro",
                    "安全简介": "safe_intro",
                    "topic": "topic",
                    "题材": "topic",
                    "level": "level",
                    "难度": "level",
                }[metadata[1].casefold()]
                current[field] = metadata[2]
                block["role"] = "metadata:" + field
            elif mode in {"word", "chunk"}:
                block["role"] = "attachment:" + mode
                current["attachments"].append(
                    dict(block) | {"entry_type": mode, "source_locations": [block["id"]]}
                )
            else:
                lang = (
                    mode
                    if mode in {"english", "chinese"}
                    else ("chinese" if re.search(r"[\u4e00-\u9fff]", text) else "english")
                )
                block["role"] = lang
                current[lang].append(block)
        elif isinstance(part, Table):
            table_no += 1
            if current is None:
                current = new_article()
            table_rows = []
            for row_no, row in enumerate(part.rows, start=1):
                cells = []
                for column_no, cell in enumerate(row.cells, start=1):
                    block = {
                        "id": f"table-{table_no:04d}-row-{row_no:04d}-cell-{column_no:04d}",
                        "kind": "table_cell",
                        "table_index": table_no,
                        "row_index": row_no,
                        "column_index": column_no,
                        "text": cell.text,
                        "role": "attachment:" + mode if mode in {"word", "chunk"} else "unconsumed",
                    }
                    block["children"] = original_inner_blocks(
                        cell.iter_inner_content(), block["id"]
                    )
                    current["blocks"].append(block)
                    cells.append(block)
                table_rows.append(cells)
            headers = (
                [LEXICON_COLUMNS.get(c["text"].strip().casefold(), "") for c in table_rows[0]]
                if table_rows
                else []
            )
            supported = (
                mode in {"word", "chunk"}
                and "english" in headers
                and "meaning_zh" in headers
                and len([h for h in headers if h]) == len(set(headers) - {""})
            )
            for index, cells in enumerate(table_rows):
                if supported and index == 0:
                    for cell_block in cells:
                        cell_block["role"] = "lexicon_table_header"
                    continue
                locations = [c["id"] for c in cells]
                if mode in {"word", "chunk"}:
                    current["attachments"].append(
                        {
                            "kind": "table_row",
                            "text": "\t".join(c["text"] for c in cells),
                            "entry_type": mode,
                            "source_locations": locations,
                            "fields": {
                                h: c["text"].strip()
                                for h, c in zip(headers, cells, strict=False)
                                if h
                            }
                            if supported
                            and not any(
                                child["kind"] == "table" for c in cells for child in c["children"]
                            )
                            else None,
                        }
                    )
                else:
                    current["issues"].append(
                        parsing_issue(
                            "IMPORT_TABLE_UNSUPPORTED",
                            "document_path",
                            "正文表格保留原输入,需人工整理对应关系",
                            locations,
                        )
                    )
    results: list[dict[str, Any]] = []
    for article in articles:
        candidate = {
            k: article[k]
            for k in ("title_en", "title_zh", "source", "copyright", "safe_intro", "topic", "level")
            if k in article
        }
        if not article["english"] or not article["chinese"]:
            article["issues"].append(
                parsing_issue(
                    "BILINGUAL_PAIRING_REQUIRED",
                    "document_path",
                    "正文须含完整英文和中文",
                    [b["id"] for b in article["blocks"]],
                )
            )
        if kind == "short_reading":
            candidate.update(
                paragraphs_en=[b["text"].strip() for b in article["english"]],
                translation_zh="\n\n".join(b["text"].strip() for b in article["chinese"]),
            )
        else:
            article["dialogue"] = pair_dialogue(article)
            candidate.update(
                schema_version=2, dialogue=article["dialogue"], vocabulary=[], chunks=[]
            )
        entries = materialize_docx_entries(article, kind)
        snapshot = {
            "schema_version": 1,
            "origin": "docx",
            "document_sha256": hashlib.sha256(data).hexdigest(),
            "article_index": len(results) + 1,
            "title": {k: article[k] for k in ("title_en", "title_zh")},
            "blocks": deepcopy(article["blocks"]),
            "attachments": deepcopy(article["attachments"]),
            "conflicts": {},
        }
        candidate.update(
            source_snapshot=snapshot, parsed_entries=entries, parsing_issues=article["issues"]
        )
        results.append(candidate)
    return results


def parse_package(
    data: bytes, kind: ContentType, *, limits: ImportLimits | None = None
) -> dict[str, Any]:
    limits = limits or ImportLimits()
    files = read_container(data, limits, spool=True)
    try:
        return parse_package_files(data, kind, limits, files)
    finally:
        if isinstance(files, SpoolFiles):
            files.close()


def parse_package_files(
    data: bytes,
    kind: ContentType,
    limits: ImportLimits,
    files: Mapping[str, bytes],
) -> dict[str, Any]:
    content_table = files.get("内容表.xlsx") or files.get("content.xlsx")
    if not content_table:
        raise AppError("IMPORT_CONTENT_TABLE_REQUIRED", "缺少内容表.xlsx", 422)
    original_rows = read_sheet(content_table, limits, include_empty=True)
    rows = [
        row
        for row in original_rows
        if any(
            value
            for field, value in row.items()
            if field not in {"row_no", "formula_columns", "original_row"}
        )
    ]
    if len(rows) > limits.max_items:
        raise AppError("IMPORT_LIMIT_EXCEEDED", "单逻辑批次最多1000篇", 422)
    keys = Counter(row.get("import_content_id", "") for row in rows)
    original_word_rows = (
        read_sheet(files["词汇表.xlsx"], limits, include_empty=True)
        if "词汇表.xlsx" in files
        else []
    )
    word_rows = [
        row
        for row in original_word_rows
        if any(
            value
            for field, value in row.items()
            if field not in {"row_no", "formula_columns", "original_row"}
        )
    ]
    word_keys = Counter(
        (r.get("import_content_id", ""), r.get("import_entry_id") or r.get("entry_id", ""))
        for r in word_rows
    )
    words_by_parent: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for word in word_rows:
        words_by_parent[word.get("import_content_id", "")].append(word)
    referenced = {"内容表.xlsx", "content.xlsx", "词汇表.xlsx"}
    items = []
    for row in rows:
        issues: list[dict[str, Any]] = []

        def issue(
            code: str,
            field: str,
            message: str,
            severity: str = "ERROR",
            *,
            target: list[dict[str, Any]] = issues,
            row_no: int = row["row_no"],
        ) -> None:
            target.append(
                {
                    "code": code,
                    "field": field,
                    "message": message,
                    "severity": severity,
                    "row_no": row_no,
                }
            )

        key = row.get("import_content_id", "")
        if not key or len(key) > 64:
            issue("IMPORT_ID_REQUIRED", "import_content_id", "填写稳定内容导入ID,长度不超过64")
        if key and keys[key] > 1:
            issue("IMPORT_DUPLICATE_ID", "import_content_id", "同批内容ID重复,请修正或排除")
        if row["formula_columns"]:
            issue("IMPORT_FORMULA_INVALID", "formula_columns", "公式不计算,请转为明确字面值")
        bound: dict[str, dict[str, str]] = {}
        for field in ("document_path", "cover_path", "learning_image_path", "audio_path"):
            path = row.get(field, "")
            if not path:
                if field == "document_path":
                    issue("IMPORT_DOCUMENT_REQUIRED", field, "每篇需一个双语DOCX")
                elif field == "audio_path" and kind == "dialogue":
                    issue("AUDIO_MISSING", field, "场景发布前需完整对话音频", "DRAFT")
                continue
            path = safe_path(path, limits)
            referenced.add(path)
            ext = path.rsplit(".", 1)[-1].casefold()
            allowed = {
                "document_path": {"docx"},
                "cover_path": {"png", "jpg", "jpeg", "webp"},
                "learning_image_path": {"png", "jpg", "jpeg", "webp"},
                "audio_path": {"mp3", "m4a", "wav", "aac"},
            }[field]
            if path not in files or ext not in allowed:
                issue("IMPORT_RESOURCE_INVALID", field, "路径不存在或文件类型不符")
            elif kind == "short_reading" and field == "learning_image_path":
                issue("IMPORT_ROLE_INVALID", field, "纯阅读短文不配置受限学习原图")
            else:
                bound[field] = {
                    "path": path,
                    "sha256": hashlib.sha256(files[path]).hexdigest(),
                    "role": field,
                }
        candidate: dict[str, Any] = {}
        documents: list[dict[str, Any]] = []
        if "document_path" in bound:
            try:
                articles = parse_docx(files[bound["document_path"]["path"]], kind, limits)
                documents = [a["source_snapshot"] for a in articles]
                if len(articles) != 1:
                    issue(
                        "IMPORT_DOCUMENT_BOUNDARY_INVALID",
                        "document_path",
                        "统一ZIP每行DOCX必须恰好一篇",
                    )
                elif articles:
                    candidate = articles[0]
                    issues.extend(candidate["parsing_issues"])
            except AppError as error:
                issue(error.code, "document_path", "文档损坏或包含不安全资料,请重新保存")
        source_snapshot = deepcopy(
            candidate.get(
                "source_snapshot",
                {
                    "schema_version": 1,
                    "origin": "zip",
                    "title": {},
                    "blocks": [],
                    "attachments": [],
                    "conflicts": {},
                },
            )
        )
        source_snapshot.update(
            origin="zip",
            document_path=row.get("document_path", ""),
            workbook={
                "content_row": deepcopy(row["original_row"]),
                "entry_rows": [deepcopy(w["original_row"]) for w in words_by_parent.get(key, [])],
            },
            resources=deepcopy(bound),
        )
        if len(documents) > 1:
            source_snapshot["documents"] = deepcopy(documents)
        for field in (
            "title_en",
            "title_zh",
            "category",
            "safe_intro",
            "topic",
            "level",
            "source_copyright",
        ):
            value = row.get(field, "")
            if not value:
                continue
            if (
                field in {"title_en", "title_zh"}
                and candidate.get(field)
                and value != candidate[field]
            ):
                issue("IMPORT_TEXT_CONFLICT", field, "表格与文档标题不同,需明确确认", "CONFIRM")
                source_snapshot["conflicts"][field] = {
                    "document": candidate[field],
                    "workbook": value,
                }
            candidate[field if field != "source_copyright" else "copyright"] = value
        if kind == "short_reading":
            candidate.setdefault("category", "")
            for field in ("safe_intro", "topic", "level"):
                candidate.setdefault(field, "")
        elif candidate.get("dialogue"):
            if any(not s["chinese"] for s in candidate["dialogue"]):
                issue(
                    "BILINGUAL_PAIRING_REQUIRED",
                    "document_path",
                    "缺少对应中文或配对不清,需校对",
                    "DRAFT",
                )
        entries = []
        for word in words_by_parent.get(key, []) if kind != "short_reading" else []:
            if word.get("import_content_id") != key:
                continue
            word_key = word.get("import_entry_id") or word.get("entry_id", "")
            item = {
                k: v
                for k, v in word.items()
                if k not in {"formula_columns", "row_no", "original_row"}
            }
            item["import_entry_id"] = word_key
            if not word_key or word_keys[(key, word_key)] > 1:
                issue("IMPORT_DUPLICATE_ENTRY_ID", "import_entry_id", "同来源条目ID为空或重复")
            if word["formula_columns"]:
                issue("IMPORT_FORMULA_INVALID", "import_entry_id", "条目公式需转为字面值")
            if word.get("entry_type") not in {"word", "chunk"}:
                issue("ENTRY_TYPE_INVALID", "entry_type", "条目类型须为word或chunk")
            if word.get("entry_type") == "chunk" and word.get("image_path"):
                issue("CHUNK_IMAGE_INVALID", "image_path", "语块不接收图片")
            for field in ("image_path", "pronunciation_path"):
                if not word.get(field):
                    continue
                path = safe_path(word[field], limits)
                referenced.add(path)
                if path not in files:
                    issue("IMPORT_RESOURCE_INVALID", field, "条目素材路径不存在")
                else:
                    allowed = (
                        {".png", ".jpg", ".jpeg", ".webp"}
                        if field == "image_path"
                        else {".mp3", ".wav", ".m4a", ".aac"}
                    )
                    if PurePosixPath(path).suffix.lower() not in allowed:
                        issue(
                            "IMPORT_RESOURCE_EXTENSION_INVALID", field, "素材扩展名不符合角色要求"
                        )
                    item[field + "_hash"] = hashlib.sha256(files[path]).hexdigest()
            if kind == "dialogue":
                resolved_source = resolve_entry_source(item, candidate.get("dialogue", []))
                for source_issue in resolved_source["issues"]:
                    issue(source_issue["code"], "source_locator", source_issue["message"], "DRAFT")
                item.update(
                    {
                        name: resolved_source[name]
                        for name in ("source_sentence", "source_sentence_ids", "source_locator")
                    }
                )
            elif isinstance(item.get("source_locator"), str) and item["source_locator"]:
                try:
                    locator = json.loads(item["source_locator"])
                except (ValueError, TypeError):
                    locator = None
                item["source_locator"] = (
                    {str(k): str(v) for k, v in locator.items()}
                    if isinstance(locator, dict)
                    else {"reference": item["source_locator"]}
                )
            entries.append(item)
        docx_entries = candidate.get("parsed_entries", [])
        if kind != "short_reading":
            if words_by_parent.get(key) and docx_entries:
                source_snapshot["conflicts"]["lexicon"] = {
                    "document": deepcopy(docx_entries),
                    "workbook": deepcopy(entries),
                }
                issue(
                    "IMPORT_LEXICON_CONFLICT",
                    "entries",
                    "同时提供文档词表和Excel词表,保留两侧供确认",
                    "CONFIRM",
                )
            elif not words_by_parent.get(key):
                entries = deepcopy(docx_entries)
                for entry in entries:
                    for field in ("image_path", "pronunciation_path"):
                        if entry.get(field):
                            path = safe_path(entry[field], limits)
                            referenced.add(path)
                            if path not in files:
                                issue("IMPORT_RESOURCE_INVALID", field, "DOCX条目素材路径不存在")
                            else:
                                allowed = (
                                    {".png", ".jpg", ".jpeg", ".webp"}
                                    if field == "image_path"
                                    else {".mp3", ".wav", ".m4a", ".aac"}
                                )
                                if PurePosixPath(path).suffix.lower() not in allowed:
                                    issue(
                                        "IMPORT_RESOURCE_EXTENSION_INVALID",
                                        field,
                                        "素材扩展名不符合角色要求",
                                    )
                                entry[field + "_hash"] = hashlib.sha256(files[path]).hexdigest()
        candidate["source_snapshot"] = deepcopy(source_snapshot)
        defaults = ["default_cover"] if not row.get("cover_path") else []
        defaults += [
            "no_image:" + e["import_entry_id"]
            for e in entries
            if e.get("entry_type") == "word" and not e.get("image_path")
        ]
        input_hash = canonical_hash(
            {
                "rules": "v15.1",
                "type": kind,
                "candidate": semantic_input(candidate),
                "resources": bound,
                "entries": semantic_input(entries),
            }
        )
        items.append(
            {
                "row_no": row["row_no"],
                "import_content_id": key,
                "candidate": candidate,
                "source_snapshot": source_snapshot,
                "entries": entries,
                "resources": bound,
                "issues": issues,
                "defaults": defaults,
                "input_hash": input_hash,
                "expected_content_version": row.get("expected_content_version") or None,
            }
        )
    orphan_rows = [
        r
        for r in word_rows
        if not r.get("import_content_id") or r.get("import_content_id") not in keys
    ]
    orphans = [r["row_no"] for r in orphan_rows]
    batch_issues = [
        {
            "code": "IMPORT_ORPHAN_ENTRY",
            "severity": "ERROR",
            "row_no": r["row_no"],
            "field": "import_content_id",
            "import_content_id": r.get("import_content_id", ""),
            "message": "词汇表来源内容ID为空或不在内容表,请修正后重新上传",
        }
        for r in orphan_rows
    ]
    if kind == "short_reading" and word_rows:
        batch_issues.append(
            {
                "code": "IMPORT_READING_ENTRIES_FORBIDDEN",
                "severity": "ERROR",
                "field": "词汇表.xlsx",
                "message": "纯阅读短文不接收词条附件",
            }
        )
    return {
        "namespace": IMPORT_NAMESPACE,
        "parser_version": "v15.1",
        "content_type": kind,
        "original_sha256": hashlib.sha256(data).hexdigest(),
        "source_snapshot": {
            "schema_version": 1,
            "origin": "zip",
            "workbook": {
                "content_rows": [r["original_row"] for r in original_rows],
                "entry_rows": [r["original_row"] for r in original_word_rows],
            },
            "files": sorted(files),
        },
        "items": items,
        "unreferenced_files": sorted(set(files) - referenced),
        "orphan_entry_rows": orphans,
        "issues": batch_issues,
        "limits": asdict(limits),
    }
