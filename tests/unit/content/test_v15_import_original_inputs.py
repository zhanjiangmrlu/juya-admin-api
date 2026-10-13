"""Regressions for input loss, guessed pairing and supplied lexicon loss."""

import json
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from docx import Document

from juya_admin_api.modules.content_import.parser import ImportLimits, parse_docx, parse_package
from juya_admin_api.modules.content_import.service import ImportService
from juya_admin_api.modules.content_import.stages import Claim
from tests.unit.content.test_v15_import_parser import docx_bytes, package, xlsx_bytes
from tests.unit.content.test_v15_import_service import service_fixture


def document(lines, table=None):
    doc = Document()
    doc.add_heading("你好\uff5cHello", level=1)
    for line in lines:
        doc.add_paragraph(line)
    if table:
        target = doc.add_table(rows=0, cols=len(table[0]))
        for values in table:
            for cell, value in zip(target.add_row().cells, values, strict=True):
                cell.text = value
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def article(lines, kind="dialogue", table=None):
    return parse_docx(document(lines, table), kind, ImportLimits())[0]


@pytest.mark.parametrize(
    "lines",
    [
        ["A: Hello.", "A\uff1a你好。", "B: Goodbye.", "B\uff1a再见。"],
        ["English", "A: Hello.", "B: Goodbye.", "Chinese", "A\uff1a你好。", "B\uff1a再见。"],
        ["A: Hello.", "甲\uff1a你好。", "B: Goodbye.", "乙\uff1a再见。"],
    ],
)
def test_original_order_speakers_and_locations_survive_both_supported_pairing_layouts(lines):
    result = article(lines)
    assert [s["english"] for s in result["dialogue"]] == ["Hello.", "Goodbye."]
    assert [s["chinese"] for s in result["dialogue"]] == ["你好。", "再见。"]
    assert not result["parsing_issues"]
    original = result["source_snapshot"]
    assert original["title"] == {"title_zh": "你好", "title_en": "Hello"}
    assert [p["text"] for p in original["blocks"]] == ["你好\uff5cHello", *lines]
    assert original["blocks"][1]["id"] == "paragraph-0002"
    expected_location = "paragraph-0003" if lines[0] == "English" else "paragraph-0002"
    assert result["dialogue"][0]["source_locations"][0] == expected_location


def test_alternating_missing_and_extra_translation_do_not_shift_other_turns():
    result = article(["A: Hello.", "B: Goodbye.", "B\uff1a再见。", "C\uff1a多余译文。"])
    assert [s["chinese"] for s in result["dialogue"]] == ["", "再见。"]
    issues = result["parsing_issues"]
    assert {i["code"] for i in issues} >= {"BILINGUAL_PAIRING_REQUIRED", "BILINGUAL_EXTRA_CHINESE"}
    assert all(i["source_locations"] for i in issues)
    assert result["source_snapshot"]["blocks"][-1]["text"] == "C\uff1a多余译文。"


def test_reversed_equal_speaker_names_are_candidates_with_explicit_conflict():
    result = article(
        ["English", "A: Hello.", "B: Goodbye.", "Chinese", "B\uff1a再见。", "A\uff1a你好。"]
    )
    assert [s["speaker"] for s in result["dialogue"]] == ["A", "B"]
    assert [s["chinese_speaker"] for s in result["dialogue"]] == ["B", "A"]
    assert sum(i["code"] == "BILINGUAL_SPEAKER_CONFLICT" for i in result["parsing_issues"]) == 2


def test_grouped_count_difference_and_inconsistent_alias_sequence_require_proofreading():
    result = article(
        [
            "English",
            "A: Hello.",
            "A: Again.",
            "Chinese",
            "甲\uff1a你好。",
            "乙\uff1a再说。",
            "甲\uff1a多余。",
        ]
    )
    assert {i["code"] for i in result["parsing_issues"]} >= {
        "BILINGUAL_SPEAKER_CONFLICT",
        "BILINGUAL_EXTRA_CHINESE",
    }


def test_supplied_line_lexicon_metadata_and_unparsed_text_are_preserved():
    result = article(
        [
            "A: Hello, neighbour.",
            "甲\uff1a你好\uff0c邻居。",
            "",
            "Vocabulary",
            "neighbour | 邻居 | 同住附近的人 | Hello, neighbour.",
            "unstructured attachment",
            "Useful Chunks",
            "Hello, neighbour | 你好邻居",
            "Source: User book",
            "Copyright: Licensed",
            "Safe Intro: Friendly greeting",
        ]
    )
    word, chunk = result["parsed_entries"]
    assert (word["english"], word["meaning_zh"], word["explanation"]) == (
        "neighbour",
        "邻居",
        "同住附近的人",
    )
    assert word["source_sentence_ids"] == ["sentence-0001"]
    assert word["import_entry_id"].startswith("docx-word-")
    assert chunk["entry_type"] == "chunk" and chunk["english"] == "Hello, neighbour"
    assert result["source"] == "User book" and result["copyright"] == "Licensed"
    assert result["safe_intro"] == "Friendly greeting"
    assert any(i["code"] == "IMPORT_LEXICON_AMBIGUOUS" for i in result["parsing_issues"])
    assert "unstructured attachment" in [p["text"] for p in result["source_snapshot"]["blocks"]]
    assert "" in [p["text"] for p in result["source_snapshot"]["blocks"]]


def test_explicit_lexicon_table_keeps_cell_positions_and_supplied_identifiers():
    result = article(
        ["A: Hello, neighbour.", "甲\uff1a你好\uff0c邻居。", "Vocabulary"],
        table=[
            [
                "import_entry_id",
                "english",
                "meaning_zh",
                "phonetic",
                "explanation",
                "source_sentence",
            ],
            [
                "provided-1",
                "neighbour",
                "邻居",
                "/\u02c8ne\u026abə/",
                "附近的人",
                "Hello, neighbour.",
            ],
            ["unclear", "bad", "", "", "", ""],
        ],
    )
    entry = result["parsed_entries"][0]
    assert entry["import_entry_id"] == "provided-1" and entry["phonetic"] == "/\u02c8ne\u026abə/"
    assert entry["source_sentence_ids"] == ["sentence-0001"]
    assert any(
        b["id"] == "table-0001-row-0003-cell-0002" and b["text"] == "bad"
        for b in result["source_snapshot"]["blocks"]
    )
    assert any(i["code"] == "IMPORT_LEXICON_AMBIGUOUS" for i in result["parsing_issues"])


def test_short_reading_ignores_lexicon_but_retains_original_lists():
    result = article(
        ["Hello.", "你好。", "Vocabulary", "hello | 你好", "Useful Chunks", "say hello | 打招呼"],
        "short_reading",
    )
    assert result["paragraphs_en"] == ["Hello."] and result["translation_zh"] == "你好。"
    assert result["parsed_entries"] == [] and "vocabulary" not in result
    assert len(result["source_snapshot"]["attachments"]) == 2


def test_zip_conflicting_titles_and_lists_preserve_both_inputs_and_workbook_priority():
    data = package(
        {
            "内容表.xlsx": xlsx_bytes(
                [
                    ["import_content_id", "document_path", "title_en", "title_zh", "custom"],
                    ["one", "one.docx", "Workbook title", "表格标题", "  supplied text  "],
                ]
            ),
            "one.docx": document(
                [
                    "A: Hello, neighbour.",
                    "甲\uff1a你好\uff0c邻居。",
                    "Vocabulary",
                    "neighbour | 文档释义",
                ]
            ),
            "词汇表.xlsx": xlsx_bytes(
                [
                    [
                        "import_content_id",
                        "import_entry_id",
                        "entry_type",
                        "english",
                        "meaning_zh",
                        "source_sentence",
                    ],
                    ["one", "word-1", "word", "neighbour", "表格释义", "Hello, neighbour."],
                ]
            ),
        }
    )
    result = parse_package(data, "dialogue")
    item = result["items"][0]
    assert item["candidate"]["title_en"] == "Workbook title"
    snapshot = item["source_snapshot"]
    assert snapshot["title"]["title_en"] == "Hello"
    assert snapshot["workbook"]["content_row"]["values"]["custom"] == "  supplied text  "
    assert snapshot["conflicts"]["title_en"] == {"document": "Hello", "workbook": "Workbook title"}
    assert snapshot["attachments"][0]["text"] == "neighbour | 文档释义"
    assert item["entries"][0]["meaning_zh"] == "表格释义"
    assert any(i["code"] == "IMPORT_LEXICON_CONFLICT" for i in item["issues"])
    item["candidate"]["source_snapshot"]["title"]["title_en"] = "edited"
    assert snapshot["title"]["title_en"] == "Hello"


def test_zip_without_workbook_lexicon_materializes_docx_entries():
    result = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "one.docx"]]
                ),
                "one.docx": document(["Hello.", "你好。", "Vocabulary", "hello | 你好"]),
            }
        ),
        "word_card",
    )
    assert result["items"][0]["entries"][0]["english"] == "hello"


@pytest.mark.asyncio
async def test_docx_validate_persists_original_lexicon_and_diagnostics():
    service, _ = service_fixture()
    batch = service.batch.return_value
    batch.update(import_mode="DOCX", status="PENDING", input_hash="source-hash")
    service.original.return_value = document(["Hello.", "Vocabulary", "hello | 你好"])
    service.stages = AsyncMock()
    service.stages.claim.return_value = Claim("stage", 1, "source-hash", "v15.1", "worker")
    service.assert_claim_within = AsyncMock()
    await ImportService.validate(service, "batch", "stage")
    calls = service.sessions.return_value.__aenter__.return_value.execute.await_args_list
    inserted = next(
        call.args[1] for call in calls if "INSERT INTO content_import_item" in str(call.args[0])
    )
    item = json.loads(inserted["candidate"])
    assert item["entries"][0]["english"] == "hello"
    assert item["source_snapshot"]["title"]["title_en"] == "Hello"
    assert any(i["code"] == "BILINGUAL_PAIRING_REQUIRED" for i in json.loads(inserted["issues"]))


def test_existing_alias_fixture_remains_accepted_and_stable():
    original_document = docx_bytes()
    one = parse_docx(original_document, "dialogue", ImportLimits())[0]
    two = parse_docx(original_document, "dialogue", ImportLimits())[0]
    assert not one["parsing_issues"]
    assert one["dialogue"][0]["speaker"] == "A"
    assert one["dialogue"][0]["chinese_speaker"] == "甲"
    assert one["source_snapshot"] == two["source_snapshot"]


@pytest.mark.parametrize(
    "chinese", [["甲\uff1a你好。", "乙\uff1a再见。"], ["B\uff1a再见。", "A\uff1a你好。"]]
)
def test_grouped_without_section_headings_keeps_order_and_diagnoses_reversed_names(chinese):
    result = article(["A: Hello.", "B: Goodbye.", *chinese])
    assert len(result["dialogue"][0]["source_locations"]) == 2
    assert result["dialogue"][0]["chinese_speaker"] == chinese[0].split("\uff1a")[0]
    if chinese[0].startswith("B"):
        assert sum(i["code"] == "BILINGUAL_SPEAKER_CONFLICT" for i in result["parsing_issues"]) == 2


def test_workbook_reordering_does_not_change_identity_or_semantic_input_hash():
    original_document = docx_bytes()

    def build(keys):
        return parse_package(
            package(
                {
                    "内容表.xlsx": xlsx_bytes(
                        [
                            ["import_content_id", "document_path"],
                            *[[key, "one.docx"] for key in keys],
                        ]
                    ),
                    "one.docx": original_document,
                }
            ),
            "word_card",
        )

    first, second = build(["one", "two"]), build(["two", "one"])
    assert first["items"][0]["import_content_id"] == second["items"][1]["import_content_id"]
    assert first["items"][0]["input_hash"] == second["items"][1]["input_hash"]


def test_blank_and_orphan_workbook_rows_are_retained_in_batch_original():
    result = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "document_path"],
                        ["one", "one.docx"],
                        ["", ""],
                        ["two", "one.docx"],
                    ]
                ),
                "one.docx": docx_bytes(),
                "词汇表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "import_entry_id"],
                        ["missing", "unmatched"],
                        ["", ""],
                        ["one", "word"],
                    ]
                ),
            }
        ),
        "word_card",
    )
    original = result["source_snapshot"]["workbook"]
    assert [r["row_no"] for r in original["content_rows"]] == [2, 3, 4]
    assert [r["row_no"] for r in original["entry_rows"]] == [2, 3, 4]
    assert original["entry_rows"][0]["values"]["import_entry_id"] == "unmatched"


def test_ambiguous_multiarticle_zip_preserves_every_article_original():
    doc = Document()
    for title in ["第一篇\uff5cFirst", "第二篇\uff5cSecond"]:
        doc.add_heading(title, level=1)
        doc.add_paragraph("Hello.")
        doc.add_paragraph("你好。")
    output = BytesIO()
    doc.save(output)
    result = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "one.docx"]]
                ),
                "one.docx": output.getvalue(),
            }
        ),
        "dialogue",
    )
    item = result["items"][0]
    assert len(item["source_snapshot"]["documents"]) == 2
    assert item["source_snapshot"]["documents"][1]["blocks"][0]["id"] == "paragraph-0004"
    assert any(i["code"] == "IMPORT_DOCUMENT_BOUNDARY_INVALID" for i in item["issues"])


@pytest.mark.asyncio
async def test_save_draft_strips_parse_metadata_but_keeps_import_original():
    from unittest.mock import Mock

    from juya_admin_api.modules.content.schemas import SceneContent

    service, _ = service_fixture()
    candidate = article(["A: Hello.", "甲\uff1a你好。"])
    item = {
        "operation": "CREATE",
        "candidate": candidate,
        "entries": [],
        "source_snapshot": json.loads(json.dumps(candidate["source_snapshot"])),
        "import_content_id": "one",
        "input_hash": "input-hash",
    }
    service.production.template.return_value = 1
    service.production.save_revision_within.side_effect = lambda session, revision, expected: (
        revision
    )
    session = AsyncMock()
    session.scalar.return_value = 1
    empty = Mock()
    empty.first.return_value = None
    session.execute.return_value = empty
    await ImportService.save_item_within(
        service, session, {"content_type": "dialogue", "series_id": "series"}, item, None, "actor"
    )
    revision = service.production.save_revision_within.await_args.args[1]
    SceneContent.model_validate(revision.content)
    assert item["source_snapshot"]["blocks"][2]["text"] == "甲\uff1a你好。"


def test_duplicate_supplied_entries_remain_original_without_colliding_canonical_keys():
    result = article(
        [
            "Hello.",
            "你好。",
            "Vocabulary",
            "hello | 你好",
            "hello | 你好",
        ]
    )
    assert result["parsed_entries"] == []
    assert len(result["source_snapshot"]["attachments"]) == 2
    assert sum(i["code"] == "IMPORT_DUPLICATE_ENTRY_ID" for i in result["parsing_issues"]) == 2


def test_lexicon_table_type_conflict_is_not_silently_overwritten():
    result = article(
        ["Hello.", "你好。", "Vocabulary"],
        table=[
            ["english", "meaning_zh", "entry_type"],
            ["hello there", "你好呀", "chunk"],
        ],
    )
    assert result["parsed_entries"] == []
    assert any(i["code"] == "ENTRY_TYPE_INVALID" for i in result["parsing_issues"])
    assert result["source_snapshot"]["attachments"][0]["fields"]["entry_type"] == "chunk"


def test_docx_entry_media_extension_obeys_same_zip_resource_rules_as_workbook():
    result = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "one.docx"]]
                ),
                "one.docx": document(
                    ["Hello.", "你好。", "Vocabulary"],
                    table=[
                        ["english", "meaning_zh", "image_path"],
                        ["hello", "你好", "images/a.txt"],
                    ],
                ),
                "images/a.txt": b"wrong role",
            }
        ),
        "word_card",
    )
    assert any(
        i["code"] == "IMPORT_RESOURCE_EXTENSION_INVALID" for i in result["items"][0]["issues"]
    )


def test_unchanged_legacy_zip_replays_with_its_existing_semantic_hash():
    from juya_admin_api.modules.content_import.parser import canonical_hash

    item = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "one.docx"]]
                ),
                "one.docx": docx_bytes(),
            }
        ),
        "dialogue",
    )["items"][0]
    old_candidate = {
        "title_en": "Hello",
        "title_zh": "你好",
        "source": "",
        "copyright": "",
        "schema_version": 2,
        "dialogue": [
            {
                "id": "sentence-0001",
                "speaker": "A",
                "english": "Hello, neighbour.",
                "chinese": "你好\uff0c邻居。",
            }
        ],
        "vocabulary": [],
        "chunks": [],
    }
    assert item["input_hash"] == canonical_hash(
        {
            "rules": "v15.1",
            "type": "dialogue",
            "candidate": old_candidate,
            "resources": item["resources"],
            "entries": [],
        }
    )


def test_unsupported_nested_tables_keep_all_original_cell_paragraphs():
    doc = Document()
    doc.add_heading("你好\uff5cHello", level=1)
    doc.add_paragraph("Hello.")
    doc.add_paragraph("你好。")
    doc.add_paragraph("Vocabulary")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ["english", "meaning_zh"], strict=True):
        cell.text = text
    table.cell(1, 0).text = "hello"
    table.cell(1, 1).text = "你好"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "nested original explanation"
    output = BytesIO()
    doc.save(output)
    result = parse_docx(output.getvalue(), "dialogue", ImportLimits())[0]
    block = next(
        b for b in result["source_snapshot"]["blocks"] if b["id"] == "table-0001-row-0002-cell-0002"
    )
    nested_original = block["children"][1]["rows"][0][0]["children"][0]
    assert nested_original["text"] == "nested original explanation"
    assert result["parsed_entries"] == []
    assert any(i["code"] == "IMPORT_LEXICON_AMBIGUOUS" for i in result["parsing_issues"])


@pytest.mark.asyncio
async def test_unchanged_legacy_docx_keeps_existing_content_key_and_input_hash():
    from juya_admin_api.modules.content_import.parser import canonical_hash

    service, _ = service_fixture()
    batch = service.batch.return_value
    batch.update(import_mode="DOCX", status="PENDING", input_hash="source-hash")
    service.original.return_value = docx_bytes()
    service.stages = AsyncMock()
    service.stages.claim.return_value = Claim("stage", 1, "source-hash", "v15.1", "worker")
    service.assert_claim_within = AsyncMock()
    await ImportService.validate(service, "batch", "stage")
    calls = service.sessions.return_value.__aenter__.return_value.execute.await_args_list
    inserted = next(
        call.args[1] for call in calls if "INSERT INTO content_import_item" in str(call.args[0])
    )
    item = json.loads(inserted["candidate"])
    assert item["import_content_id"] == canonical_hash(["source-hash", 0])
    old_candidate = {
        "title_en": "Hello",
        "title_zh": "你好",
        "source": "",
        "copyright": "",
        "schema_version": 2,
        "dialogue": [
            {
                "id": "sentence-0001",
                "speaker": "A",
                "english": "Hello, neighbour.",
                "chinese": "你好\uff0c邻居。",
            }
        ],
        "vocabulary": [],
        "chunks": [],
    }
    assert item["input_hash"] == canonical_hash(old_candidate)
