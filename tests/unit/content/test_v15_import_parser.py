import importlib.util
import json
import wave
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("你好\uff5cHello", level=1)
    doc.add_paragraph("A: Hello, neighbour.")
    doc.add_paragraph("甲\uff1a你好\uff0c邻居。")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def xlsx_bytes(rows: list[list[object]]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def package(files: dict[str, bytes]) -> bytes:
    out = BytesIO()
    with ZipFile(out, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return out.getvalue()


def parser():
    assert importlib.util.find_spec("juya_admin_api.modules.content_import.parser") is not None
    from juya_admin_api.modules.content_import.parser import parse_package

    return parse_package


def test_zip_turn_locator_is_verified_and_original_workbook_evidence_is_preserved() -> None:
    """Convert a complete declared turn to the existing stable-source output contract"""
    original_locator = json.dumps(
        {"turn": 1, "speaker": "A", "source_sentence": "Hello, neighbour."}
    )
    audio = BytesIO()
    with wave.open(audio, "wb") as sound:
        sound.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        sound.writeframes(b"\x00\x00" * 1600)
    result = parser()(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "document_path", "audio_path"],
                        ["one", "documents/one.docx", "audio/one.wav"],
                    ]
                ),
                "词汇表.xlsx": xlsx_bytes(
                    [
                        [
                            "import_content_id",
                            "import_entry_id",
                            "entry_type",
                            "english",
                            "meaning_zh",
                            "source_sentence",
                            "source_locator",
                        ],
                        [
                            "one",
                            "word-1",
                            "word",
                            "Hello",
                            "你好",
                            "Hello, neighbour.",
                            original_locator,
                        ],
                    ]
                ),
                "documents/one.docx": docx_bytes(),
                "audio/one.wav": audio.getvalue(),
            }
        ),
        "dialogue",
    )
    item = result["items"][0]
    assert not item["issues"]
    sentence_id = item["candidate"]["dialogue"][0]["id"]
    assert item["entries"][0]["source_locator"] == {"sentence_id": sentence_id}
    assert item["entries"][0]["source_sentence_ids"] == [sentence_id]
    assert item["entries"][0]["source_sentence"] == "Hello, neighbour."
    original_row = result["source_snapshot"]["workbook"]["entry_rows"][0]
    assert original_row["values"]["source_locator"] == original_locator


def test_real_short_reading_zip_uses_explicit_parent_keys_and_optional_audio() -> None:
    result = parser()(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "document_path", "topic", "level"],
                        ["one", "documents/one.docx", "邻里", "A2"],
                        ["two", "documents/two.docx", "邻里", "A2"],
                    ]
                ),
                "documents/one.docx": docx_bytes(),
                "documents/two.docx": docx_bytes(),
            }
        ),
        "short_reading",
    )
    assert [row["import_content_id"] for row in result["items"]] == ["one", "two"]
    assert result["items"][0]["candidate"]["paragraphs_en"] == ["A: Hello, neighbour."]
    assert result["items"][0]["candidate"]["translation_zh"] == "甲\uff1a你好\uff0c邻居。"
    assert "words" not in result["items"][0]["candidate"]


@pytest.mark.parametrize(
    "path",
    [
        "../escape",
        "/absolute",
        "C:/drive",
        "\\\\server\\share",
        "safe/../../escape",
        "safe:stream",
        "safe/./same",
    ],
)
def test_zip_rejects_unsafe_paths_before_any_parse(path: str) -> None:
    from juya_admin_api.shared.errors import AppError

    with pytest.raises(AppError) as failed:
        parser()(package({path: b"data"}), "short_reading")
    assert failed.value.code == "IMPORT_PATH_INVALID"


def test_zip_rejects_case_collision_and_actual_expansion_limit() -> None:
    from juya_admin_api.modules.content_import.parser import ImportLimits
    from juya_admin_api.shared.errors import AppError

    with pytest.raises(AppError):
        parser()(package({"Images/A.png": b"a", "images/a.PNG": b"b"}), "short_reading")
    with pytest.raises(AppError) as failed:
        parser()(
            package({"large": b"x" * 20}), "short_reading", limits=ImportLimits(max_total_bytes=10)
        )
    assert failed.value.code == "IMPORT_LIMIT_EXCEEDED"


def test_duplicate_content_ids_report_both_rows_and_formula_id_is_rejected() -> None:
    result = parser()(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "document_path"],
                        ["same", "documents/a.docx"],
                        ["same", "documents/a.docx"],
                        ["=1+1", "documents/a.docx"],
                    ]
                ),
                "documents/a.docx": docx_bytes(),
            }
        ),
        "short_reading",
    )
    assert all(
        any(i["code"] == "IMPORT_DUPLICATE_ID" for i in row["issues"])
        for row in result["items"][:2]
    )
    assert any(i["code"] == "IMPORT_FORMULA_INVALID" for i in result["items"][2]["issues"])


def test_word_keys_are_scoped_to_parent_and_chunk_image_is_error() -> None:
    content = xlsx_bytes(
        [
            ["import_content_id", "document_path"],
            ["one", "documents/a.docx"],
            ["two", "documents/a.docx"],
        ]
    )
    words = xlsx_bytes(
        [
            [
                "import_content_id",
                "import_entry_id",
                "entry_type",
                "english",
                "meaning_zh",
                "source_sentence",
                "image_path",
            ],
            ["one", "word_001", "word", "Hello", "你好", "Hello, neighbour.", ""],
            [
                "two",
                "word_001",
                "chunk",
                "Hello, neighbour",
                "你好邻居",
                "Hello, neighbour.",
                "images/a.png",
            ],
        ]
    )
    result = parser()(
        package(
            {
                "内容表.xlsx": content,
                "词汇表.xlsx": words,
                "documents/a.docx": docx_bytes(),
                "images/a.png": b"wrong",
            }
        ),
        "dialogue",
    )
    assert len(result["items"][0]["entries"]) == 1
    assert not any(i["code"] == "IMPORT_DUPLICATE_ENTRY_ID" for i in result["items"][1]["issues"])
    assert any(i["code"] == "CHUNK_IMAGE_INVALID" for i in result["items"][1]["issues"])


def test_orphan_entry_rows_are_batch_errors_with_row_and_parent_identity():
    result = parser()(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "documents/a.docx"]]
                ),
                "documents/a.docx": docx_bytes(),
                "词汇表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "import_entry_id", "entry_type"],
                        ["missing", "word-a", "word"],
                        ["", "word-b", "word"],
                    ]
                ),
            }
        ),
        "dialogue",
    )
    assert result["orphan_entry_rows"] == [2, 3]
    assert [issue["import_content_id"] for issue in result["issues"]] == ["missing", ""]
    assert all(
        issue["severity"] == "ERROR" and issue["code"] == "IMPORT_ORPHAN_ENTRY"
        for issue in result["issues"]
    )
    assert not result["items"][0]["entries"]


def test_short_reading_entry_attachment_cannot_be_silently_ignored():
    result = parser()(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [["import_content_id", "document_path"], ["one", "documents/a.docx"]]
                ),
                "documents/a.docx": docx_bytes(),
                "词汇表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "import_entry_id", "entry_type"],
                        ["one", "word-a", "word"],
                    ]
                ),
            }
        ),
        "short_reading",
    )
    assert not result["items"][0]["entries"]
    assert any(issue["code"] == "IMPORT_READING_ENTRIES_FORBIDDEN" for issue in result["issues"])
