"""Exercise diagnostics through real DOCX/ZIP parsing and public issue projection."""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document

from juya_admin_api.modules.content_import.parser import ImportLimits, parse_docx, parse_package
from juya_admin_api.modules.content_production.effective_issues import issue_id
from juya_admin_api.modules.content_production.projection import issues
from tests.unit.content.test_v15_import_parser import docx_bytes, package, xlsx_bytes


def conflict_article():
    """Produce a genuine Word table conflict with stable original block positions"""
    doc = Document()
    doc.add_heading("你好\uff5cHello", level=1)
    doc.add_paragraph("A: Hello.")
    doc.add_paragraph("甲\uff1a你好。")
    doc.add_paragraph("Vocabulary")
    table = doc.add_table(rows=0, cols=4)
    for values in [
        ["english", "meaning_zh", "source_sentence", "source_locator"],
        ["Hello", "你好", "Hello!", "1"],
        ["greeting", "问候", "Goodbye.", "1"],
    ]:
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    output = BytesIO()
    doc.save(output)
    return parse_docx(output.getvalue(), "dialogue", ImportLimits())[0]


def test_docx_diagnostics_survive_projection_with_unchanged_issue_identity():
    """The public schema retains rejected word-row evidence without replacing issue identities"""
    article = conflict_article()
    raw = article["parsing_issues"]
    displayed = issues(raw)
    assert len(displayed) == 2
    for item, original in zip(displayed, raw, strict=True):
        assert item.id == issue_id(original)
        diagnostic = item.model_dump()["source_diagnostic"]
        assert diagnostic["entry_text"] in {"Hello", "greeting"}
        assert diagnostic["actual_text"] == "Hello."
        assert diagnostic["source_locations"] == original["source_locations"]
    assert displayed[0].source_diagnostic.entry_text != displayed[1].source_diagnostic.entry_text


@pytest.mark.parametrize("entry_identity", ["word-1", 123])
def test_zip_diagnostics_retain_word_row_instead_of_content_row(entry_identity):
    """The content-row issue contract remains stable while diagnostics locate the exact Excel row"""
    parsed = parse_package(
        package(
            {
                "内容表.xlsx": xlsx_bytes(
                    [
                        ["import_content_id", "document_path"],
                        ["one", "documents/one.docx"],
                    ]
                ),
                "词汇表.xlsx": xlsx_bytes(
                    [
                        [
                            "import_content_id",
                            "import_entry_id",
                            "entry_type",
                            "english",
                            "meaning_zh",
                            "source_sentence",
                            "source_locator",
                        ],
                        ["one", entry_identity, "word", "Hello", "你好", "Hello!", "1"],
                    ]
                ),
                "documents/one.docx": docx_bytes(),
            }
        ),
        "dialogue",
    )
    item = parsed["items"][0]
    raw = next(value for value in item["issues"] if value["code"] == "SOURCE_TEXT_CONFLICT")
    assert raw["row_no"] == 2
    assert raw["source_diagnostic"]["source_locations"] == ["原始词条表 · 第 2 行"]
    assert raw["source_diagnostic"]["entry_id"] == str(entry_identity)
    assert raw["source_diagnostic"]["declared_locator"] == "1"
    assert issues([raw])[0].id == issue_id(
        {k: v for k, v in raw.items() if k != "source_diagnostic"}
    )
    legacy = {k: v for k, v in raw.items() if k != "source_diagnostic"}
    recovered = issues([legacy], source_snapshot=item["source_snapshot"])[0]
    assert recovered.source_diagnostic.entry_id == str(entry_identity)


def test_old_docx_snapshot_recovers_only_uniquely_attributed_diagnostics_read_only():
    """Old problem records gain verified evidence from immutable originals without being changed"""
    article = conflict_article()
    old = [
        {k: v for k, v in issue.items() if k != "source_diagnostic"}
        for issue in article["parsing_issues"]
    ]
    original = deepcopy((old, article["source_snapshot"]))
    projected = issues(old, source_snapshot=article["source_snapshot"])
    assert [issue.source_diagnostic.entry_text for issue in projected] == ["Hello", "greeting"]
    assert [issue.id for issue in projected] == [issue_id(row) for row in old]
    assert (old, article["source_snapshot"]) == original
    assert all(issue.source_diagnostic is None for issue in issues(old, source_snapshot={}))


def test_ambiguous_legacy_attribution_does_not_assign_the_first_matching_word():
    """Lacking original row evidence never makes an arbitrary repeated problem look exact"""
    article = conflict_article()
    legacy = {
        "code": "SOURCE_TEXT_CONFLICT",
        "field": "source_locator",
        "message": "来源原句与指定正文位置不一致",
        "severity": "DRAFT",
    }
    projected = issues([legacy], source_snapshot=article["source_snapshot"])
    assert projected[0].source_diagnostic is None
