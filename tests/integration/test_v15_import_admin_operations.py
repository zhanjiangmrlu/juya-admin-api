"""Private SQL, actual ZIP media bytes and fenced local-stage recovery; no cloud."""

import hashlib
import wave
from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4
from zipfile import ZipFile

import httpx
import pytest
import test_v15_worker_dispatch as dispatch
from fastapi import FastAPI
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import text
from test_v15_import_flow import import_zip

from juya_admin_api.modules.content_import.admin_stages import StageAdminStore
from juya_admin_api.modules.content_import.operations_router import create_import_operations_router
from juya_admin_api.modules.content_import.preview import ImportPreviewStore
from juya_admin_api.shared.errors import AppError

dispatched_runtime = dispatch.dispatched_runtime


def media_package():
    image = BytesIO()
    Image.new("RGB", (32, 32), (12, 34, 56)).save(image, format="PNG")
    audio = BytesIO()
    with wave.open(audio, "wb") as stream:
        stream.setnchannels(1)
        stream.setsampwidth(2)
        stream.setframerate(16000)
        stream.writeframes(b"\x10\x00" * 16000)
    with ZipFile(BytesIO(import_zip(uuid4().hex))) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    table = load_workbook(BytesIO(files["内容表.xlsx"]))
    table.active.cell(1, 4, "cover_path")
    table.active.cell(2, 4, "images/cover.png")
    table.active.cell(1, 5, "audio_path")
    table.active.cell(2, 5, "audio/full.wav")
    sheet = BytesIO()
    table.save(sheet)
    files.update(
        {
            "内容表.xlsx": sheet.getvalue(),
            "images/cover.png": image.getvalue(),
            "audio/full.wav": audio.getvalue(),
        }
    )
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return output.getvalue(), files


def operations(db):
    services = db.services
    return StageAdminStore(
        services.imports, services.authoring, services.materials, services.candidates
    )


@pytest.mark.asyncio
async def test_actual_zip_preview_credentials_bytes_and_http_ranges(dispatched_runtime):
    db = dispatched_runtime
    package, files = media_package()
    batch = await db.services.imports.create(
        series_id=db.series,
        content_type="dialogue",
        mode="ZIP",
        data=package,
        actor=db.actor,
        key=uuid4().hex,
    )
    async with db.sessions() as session:
        stage = await session.scalar(
            text(
                "SELECT public_id FROM content_import_stage WHERE item_id=:id AND stage='VALIDATE'"
            ),
            {"id": batch["batch_id"]},
        )
    await db.services.imports.validate(batch["batch_id"], stage)
    previews = ImportPreviewStore(
        db.services.imports, b"isolated-preview-secret-at-least-32", db.settings.ffprobe_path
    )
    app = FastAPI()

    async def admin():
        return SimpleNamespace(admin_user_id=db.actor)

    app.include_router(
        create_import_operations_router(
            operations(db), previews, current_admin=admin, current_admin_write=admin
        )
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        for path in ("images/cover.png", "audio/full.wav"):
            before = datetime.now(UTC)
            response = await client.post(
                f"/api/v1/admin/content-imports/{batch['batch_id']}/resources/preview",
                json={"path": path},
            )
            assert response.status_code == 200, response.text
            assert response.headers["cache-control"] == "no-store"
            issued = response.json()
            assert (
                0 < (datetime.fromisoformat(issued["expires_at"]) - before).total_seconds() <= 60.1
            )
            assert issued["sha256"] == hashlib.sha256(files[path]).hexdigest()
            assert issued["review_status"] == "ADMIN_PREVIEW_ONLY"
            token = parse_qs(urlsplit(issued["url"]).query)["token"][0]
            assert (await previews.content(batch["batch_id"], token, db.actor))[0] == files[path]
            full = await client.get(issued["url"])
            assert full.status_code == 200 and full.content == files[path]
            partial = await client.get(issued["url"], headers={"Range": "bytes=2-9"})
            assert partial.status_code == 206 and partial.content == files[path][2:10]
            assert partial.headers["content-range"] == f"bytes 2-9/{len(files[path])}"
            assert partial.headers["cache-control"] == "no-store"
            assert (
                await client.get(issued["url"], headers={"Range": "bytes=999999-"})
            ).status_code == 416
            for actor, batch_id in (("wrong-actor", batch["batch_id"]), (db.actor, "wrong-batch")):
                with pytest.raises(AppError) as denied:
                    await previews.content(batch_id, token, actor)
                assert denied.value.code == "IMPORT_PREVIEW_FORBIDDEN"
            expired = previews.credentials.issue(
                db.actor,
                batch["batch_id"],
                path,
                issued["sha256"],
                datetime.now(UTC) - timedelta(seconds=61),
            )
            with pytest.raises(AppError) as denied:
                await previews.content(batch["batch_id"], expired, db.actor)
            assert denied.value.code == "IMPORT_PREVIEW_EXPIRED"
        for path in (
            "../images/cover.png",
            "images/missing.png",
            "documents/one.docx",
            "内容表.xlsx",
        ):
            with pytest.raises(AppError) as denied:
                await previews.issue(batch["batch_id"], path, db.actor)
            assert denied.value.status_code in {404, 422}
    async with db.sessions() as session:
        original = await db.services.imports.batch(session, batch["batch_id"])
    db.services.imports.path(original["asset_id"]).write_bytes(package + b"tampered")
    with pytest.raises(AppError) as changed:
        await previews.issue(batch["batch_id"], "images/cover.png", db.actor)
    assert changed.value.code == "IMPORT_ARTIFACT_CHANGED"


@pytest.mark.asyncio
async def test_validate_manual_retry_dependency_fence_outbox_and_actual_success(dispatched_runtime):
    db = dispatched_runtime
    batch_id, stage_id = await dispatch.queued_document(db)
    stages = db.services.imports.stages
    old = await stages.claim(stage_id, "failed-validator", datetime.now(UTC))
    assert old
    await stages.fail(old, "IMPORT_FIXTURE_INTERRUPTED", datetime.now(UTC), retry=False)
    store = operations(db)
    state = await store.get(stage_id)
    payload = {
        "manual_confirmed": True,
        "expected_fence": state["fencing_version"],
        "input_hash": state["input_hash"],
        "config_version": state["config_version"],
    }
    for changes in (
        {"manual_confirmed": False},
        {"expected_fence": old.fencing_version - 1},
        {"input_hash": "0" * 64},
        {"config_version": "changed"},
    ):
        with pytest.raises(AppError):
            await store.retry(stage_id, db.actor, uuid4().hex, payload | changes)
    async with db.sessions() as session, session.begin():
        await session.execute(
            text("UPDATE content_import_batch SET status='CANCELLED' WHERE public_id=:id"),
            {"id": batch_id},
        )
    with pytest.raises(AppError) as cancelled:
        await store.retry(stage_id, db.actor, uuid4().hex, payload)
    assert cancelled.value.code == "IMPORT_CANCELLED"
    async with db.sessions() as session, session.begin():
        await session.execute(
            text("UPDATE content_import_batch SET status='PENDING' WHERE public_id=:id"),
            {"id": batch_id},
        )
    async with db.sessions() as session:
        batch = await db.services.imports.batch(session, batch_id)
    path = db.services.imports.path(batch["asset_id"])
    original = path.read_bytes()
    path.write_bytes(original + b"changed")
    with pytest.raises(AppError) as artifact:
        await store.retry(stage_id, db.actor, uuid4().hex, payload)
    assert artifact.value.code == "IMPORT_ARTIFACT_CHANGED"
    assert (await store.get(stage_id))["status"] == "FAILED"
    path.write_bytes(original)
    key = uuid4().hex
    result = await store.retry(stage_id, db.actor, key, payload)
    assert result == await store.retry(stage_id, db.actor, key, payload)
    assert result["stage_id"] == stage_id and result["fencing_version"] > old.fencing_version
    current = await store.get(stage_id)
    assert current["status"] == "QUEUED" and current["lease_until"]
    async with db.sessions() as session:
        count = await session.scalar(
            text(
                "SELECT COUNT(*) FROM admin_outbox WHERE aggregate_public_id=:id "
                "AND event_type='CONTENT_IMPORT_STAGE' AND status='PENDING'"
            ),
            {"id": stage_id},
        )
    assert count == 1
    with pytest.raises(AppError) as lost:
        await stages.finish(old, {"obsolete": True}, datetime.now(UTC))
    assert lost.value.code == "IMPORT_STAGE_LEASE_LOST"
    await db.services.imports.validate(batch_id, stage_id)
    assert (await store.get(stage_id))["status"] == "SUCCEEDED"
    assert len((await db.services.imports.preview(batch_id))["items"]) == 1
    with pytest.raises(AppError) as succeeded:
        await store.retry(stage_id, db.actor, uuid4().hex, payload)
    assert succeeded.value.code == "STAGE_RETRY_STALE"
    await db.services.imports.validate(batch_id, stage_id)
    assert (await store.get(stage_id))["attempt"] == 2
