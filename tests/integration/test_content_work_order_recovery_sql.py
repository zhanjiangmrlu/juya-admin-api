"""Counterexamples from the fresh Task 2 review, using real migrated SQL stores."""

from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy import text

from juya_admin_api.modules.content.production_store import decode, encode
from juya_admin_api.modules.content.publication_adapter import build_publication_service
from juya_admin_api.modules.content.repository import SQLAlchemyContentRepository
from juya_admin_api.modules.content.service import ContentService
from juya_admin_api.shared.errors import AppError
from tests.integration.test_content_work_orders import db as work_order_db
from tests.integration.test_content_work_orders import manual, reading

db = work_order_db


def package(count):
    document = Document()
    document.add_heading("你好|Hello", level=1)
    document.add_paragraph("A: Hello.")
    document.add_paragraph("甲:你好。")
    doc = BytesIO()
    document.save(doc)
    table = Workbook()
    table.active.append(["import_content_id", "document_path"])
    for _ in range(count):
        table.active.append([uuid4().hex, "documents/one.docx"])
    sheet = BytesIO()
    table.save(sheet)
    result = BytesIO()
    with ZipFile(result, "w") as archive:
        archive.writestr("内容表.xlsx", sheet.getvalue())
        archive.writestr("documents/one.docx", doc.getvalue())
    return result.getvalue()


async def imported_order(db, count):
    created = await db.imports.create(
        series_id=db.series,
        content_type="dialogue",
        mode="ZIP",
        data=package(count),
        actor=db.actor,
        key="upload",
    )
    async with db.sessions() as session:
        stage = await session.scalar(
            text("SELECT public_id FROM content_import_stage WHERE item_id=:id"),
            {"id": created["batch_id"]},
        )
    await db.imports.validate(created["batch_id"], stage)
    order = await db.service.from_batch(created["batch_id"], db.actor, "order")
    return order, await db.service.import_preview(order.work_order_id)


async def expire_command(db, order_id):
    async with db.sessions() as session, session.begin():
        await session.execute(
            text(
                "UPDATE content_work_command SET lease_until=UTC_TIMESTAMP(6)-INTERVAL 1 SECOND "
                "WHERE work_order_id=:id AND status='PENDING'"
            ),
            {"id": order_id},
        )


@pytest.mark.asyncio
async def test_import_retry_recovers_lost_enqueue_result_after_actual_worker_completed(
    db, monkeypatch
):
    order, preview = await imported_order(db, 1)
    await db.imports.confirm(
        preview.batch_id,
        preview.preview_hash,
        preview.preview_version,
        [preview.items[0].import_item_id],
        [],
        [],
        db.actor,
        "first-fails-default",
    )
    assert (await db.service.items(order.work_order_id)).counts.failed == 1
    queue = db.imports.queue_confirmation

    async def lost_after_worker(*args, **kwargs):
        receipt = await queue(*args, **kwargs)
        assert (await db.service.items(order.work_order_id)).counts.processing == 1
        async with db.sessions() as session:
            payload = decode(
                await session.scalar(
                    text(
                        "SELECT payload FROM admin_outbox WHERE "
                        "event_type='CONTENT_IMPORT_CONFIRM' "
                        "AND JSON_UNQUOTE(JSON_EXTRACT(payload,'$.stage_id'))=:stage"
                    ),
                    {"stage": receipt["stage_id"]},
                )
            )
        claim = await db.imports.stages.claim(
            receipt["stage_id"], "test-import-worker", datetime.now(UTC)
        )
        assert claim
        await db.imports.confirm(
            payload["batch_id"],
            **payload["confirmation"],
            actor=payload["actor"],
            key=payload["key"],
            claim=claim,
        )
        await db.imports.stages.finish(claim, {"batch_id": payload["batch_id"]}, datetime.now(UTC))
        raise RuntimeError("outer enqueue response lost after worker completed")

    monkeypatch.setattr(db.imports, "queue_confirmation", lost_after_worker)
    with pytest.raises(RuntimeError):
        await db.service.command(
            order.work_order_id,
            "import-retry",
            {
                "expected_version": order.version,
                "scope": "failed",
                "selected_item_ids": [],
                "preview_hash": preview.preview_hash,
                "preview_version": preview.preview_version,
                "acknowledged_defaults": ["default_cover"],
                "confirmed_conflicts": [],
            },
            db.actor,
            "retry",
        )
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_item WHERE public_id=:id"),
                {"id": preview.items[0].import_item_id},
            )
            == "SUCCESS"
        )
        actual_stage = await session.scalar(
            text(
                "SELECT public_id FROM content_import_stage WHERE item_id=:batch AND "
                "stage='CONFIRM'"
            ),
            {"batch": preview.batch_id},
        )
    await expire_command(db, order.work_order_id)
    monkeypatch.setattr(db.imports, "queue_confirmation", queue)
    pending = (await db.service.get(order.work_order_id)).pending_commands
    recovered = await db.service.resume(order.work_order_id, pending[0].command_id, db.actor)
    assert recovered.stage_id == actual_stage and recovered.version == order.version + 1
    assert (await db.service.get(order.work_order_id)).pending_commands == []
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM scene s JOIN content_series cs ON cs.id=s.series_id "
                    "WHERE cs.public_id=:series"
                ),
                {"series": db.series},
            )
            == 1
        )
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM content_import_stage WHERE item_id=:batch AND "
                    "stage='CONFIRM'"
                ),
                {"batch": preview.batch_id},
            )
            == 1
        )


@pytest.mark.asyncio
async def test_partial_cancel_is_authoritative_for_legacy_confirmation_and_keeps_other_row(db):
    order, preview = await imported_order(db, 2)
    first, second = (await db.service.items(order.work_order_id)).items
    await db.service.command(
        order.work_order_id,
        "cancel",
        {
            "expected_version": order.version,
            "scope": "selected",
            "selected_item_ids": [first.item_id],
        },
        db.actor,
        "cancel-one",
    )
    result = await db.imports.confirm(
        preview.batch_id,
        preview.preview_hash,
        preview.preview_version,
        [first.import_item_id, second.import_item_id],
        ["default_cover"],
        [],
        db.actor,
        "legacy-confirm",
    )
    assert result["items"][0] == {
        "item_id": first.import_item_id,
        "status": "FAILED",
        "error_code": "IMPORT_ITEM_CANCELLED",
    }
    assert result["items"][1]["status"] == "SUCCESS"
    page = await db.service.items(order.work_order_id)
    assert page.items[0].production_status == "cancelled" and page.items[0].content_id is None
    assert page.items[1].content_id and page.counts.cancelled == 1
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_item WHERE public_id=:id"),
                {"id": first.import_item_id},
            )
            == "CANCELLED"
        )
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM scene s JOIN content_series cs ON cs.id=s.series_id "
                    "WHERE cs.public_id=:series"
                ),
                {"series": db.series},
            )
            == 1
        )


@pytest.mark.asyncio
async def test_import_progress_uses_real_selected_scope_for_pending_running_and_validation(db):
    order, preview = await imported_order(db, 3)
    first, second, _third = (await db.service.items(order.work_order_id)).items
    cancel = await db.service.command(
        order.work_order_id,
        "cancel",
        {
            "expected_version": order.version,
            "scope": "selected",
            "selected_item_ids": [first.item_id],
        },
        db.actor,
        "cancel-one",
    )
    receipt = await db.service.command(
        order.work_order_id,
        "import-confirm",
        {
            "expected_version": cancel.version,
            "scope": "selected",
            "selected_item_ids": [second.item_id],
            "preview_hash": preview.preview_hash,
            "preview_version": preview.preview_version,
            "acknowledged_defaults": ["default_cover"],
            "confirmed_conflicts": [],
        },
        db.actor,
        "confirm-one",
    )
    page = await db.service.items(order.work_order_id)
    assert (
        page.counts.processing == 1
        and page.counts.review_required == 1
        and page.counts.cancelled == 1
    )
    assert page.items[1].production_status == "processing"
    assert any(t.stage_id == receipt.stage_id and t.current for t in page.items[1].tasks)
    assert not any(
        t.stage_id == receipt.stage_id and t.current
        for t in page.items[0].tasks + page.items[2].tasks
    )
    claim = await db.imports.stages.claim(
        receipt.stage_id, "actual-import-worker", datetime.now(UTC)
    )
    assert claim
    page = await db.service.items(order.work_order_id)
    assert page.items[1].production_status == "processing"
    assert (
        next(t for t in page.items[1].tasks if t.stage_id == receipt.stage_id).status == "RUNNING"
    )
    async with db.sessions() as session, session.begin():
        batch = await db.imports.batch(session, preview.batch_id, True)
        validation = await db.imports.stages.ensure_within(
            session, preview.batch_id, "VALIDATE", batch["input_hash"], "review-validation"
        )
        await session.execute(
            text(
                "INSERT INTO admin_outbox(event_id,event_type,aggregate_public_id,payload) "
                "VALUES(:id,'content.import.validate',:batch,:payload)"
            ),
            {
                "id": uuid4().hex[:26],
                "batch": preview.batch_id,
                "payload": encode({"batch_id": preview.batch_id, "stage_id": validation}),
            },
        )
        await db.imports.stages.enqueue_within(session, validation, datetime.now(UTC))
    page = await db.service.items(order.work_order_id)
    assert page.counts.processing == 2 and page.counts.cancelled == 1
    assert any(t.stage_id == validation and t.current for t in page.items[2].tasks)
    assert not any(t.current for t in page.items[0].tasks)


@pytest.mark.asyncio
@pytest.mark.parametrize("legacy_pending", [False, True])
async def test_publication_retry_recovers_exact_frozen_subset_after_worker_finishes(
    db, monkeypatch, legacy_pending
):
    publication = build_publication_service(
        db.sessions,
        db.imports.stages,
        db.production,
        ContentService(SQLAlchemyContentRepository(db.sessions)),
        db.readings,
    )
    db.service.publication = publication
    order = await manual(db)
    second = await db.readings.command(
        "create", "new", db.actor, "second", {"series_id": db.series, "content": reading("Second")}
    )
    second_item = uuid4().hex
    async with db.sessions() as session, session.begin():
        await session.execute(
            text(
                "INSERT INTO content_work_item(public_id,work_order_id,position,content_id,origin"
                "al_revision_id) VALUES(:id,:order,2,:content,:revision)"
            ),
            {
                "id": second_item,
                "order": order.work_order_id,
                "content": second["content_id"],
                "revision": second["revision_id"],
            },
        )
    items = (await db.service.items(order.work_order_id)).items
    cohort = await db.proof.create(
        [
            {
                "content_type": "short_reading",
                "content_id": i.content_id,
                "revision_id": i.revision_id,
                "candidate_hash": i.candidate_hash,
            }
            for i in items
        ],
        db.actor,
    )
    for item in items:
        await db.proof.confirm(
            cohort["public_id"],
            item.revision_id,
            item.candidate_hash,
            db.actor,
            {"text_checked": True},
        )
    await db.proof.complete(cohort["public_id"], db.actor)
    preview = await db.service.command(
        order.work_order_id,
        "publication-preview",
        {"expected_version": order.version, "scope": "all", "selected_item_ids": []},
        db.actor,
        "preview",
    )
    state = await db.service.publication_state(order.work_order_id, db.actor)
    first_receipt = await db.service.command(
        order.work_order_id,
        "publication-confirm",
        {
            "expected_version": preview.version,
            "scope": "all",
            "selected_item_ids": [],
            "publication_token": state.token,
            "range_version": state.version,
            "collection_hash": state.collection_hash,
            "manual_confirmed": True,
        },
        db.actor,
        "confirm",
    )
    publish = publication.publish_candidate

    async def fail_second(candidate, actor, key):
        if candidate["revision_id"] == second["revision_id"]:
            raise AppError("TEST_PUBLISH_RETRY_REQUIRED", "temporarily unavailable", 503)
        return await publish(candidate, actor, key)

    monkeypatch.setattr(publication, "publish_candidate", fail_second)
    await publication.run(first_receipt.stage_id, "first-publish")
    monkeypatch.setattr(publication, "publish_candidate", publish)
    state = await db.service.publication_state(order.work_order_id, db.actor)
    confirm = publication.confirm

    async def lost_after_worker(*args, **kwargs):
        receipt = await confirm(*args, **kwargs)
        await publication.run(receipt["stage_id"], "retry-publish")
        raise RuntimeError("outer retry response lost after completed publication")

    monkeypatch.setattr(publication, "confirm", lost_after_worker)
    with pytest.raises(RuntimeError):
        await db.service.command(
            order.work_order_id,
            "publication-retry",
            {
                "expected_version": first_receipt.version,
                "scope": "all",
                "selected_item_ids": [],
                "publication_token": state.token,
                "range_version": state.version,
                "collection_hash": state.collection_hash,
                "manual_confirmed": True,
            },
            db.actor,
            "retry",
        )
    state = await db.service.publication_state(order.work_order_id, db.actor)
    assert all(r.status == "SUCCEEDED" for r in state.results)
    if legacy_pending:
        async with db.sessions() as session, session.begin():
            row = (
                await session.execute(
                    text(
                        "SELECT public_id,request FROM content_work_command "
                        "WHERE work_order_id=:id AND status='PENDING' FOR UPDATE"
                    ),
                    {"id": order.work_order_id},
                )
            ).first()
            request = decode(row.request)
            request.pop("selected_revision_ids")
            first_item = next(i for i in items if i.item_id != second_item)
            request["items"].append(first_item.model_dump(mode="json"))
            await session.execute(
                text("UPDATE content_work_command SET request=:request WHERE public_id=:id"),
                {"request": encode(request), "id": row.public_id},
            )
    await expire_command(db, order.work_order_id)
    monkeypatch.setattr(publication, "confirm", confirm)
    pending = (await db.service.get(order.work_order_id)).pending_commands
    recovered = await db.service.resume(order.work_order_id, pending[0].command_id, db.actor)
    assert recovered.stage_id == state.stage_id and recovered.selected_item_ids == [second_item]
    assert recovered.version == first_receipt.version + 1
    async with db.sessions() as session:
        request = decode(
            await session.scalar(
                text("SELECT request FROM content_work_command WHERE public_id=:id"),
                {"id": recovered.command_id},
            )
        )
        assert request["selected_revision_ids"] == [second["revision_id"]]
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM content_import_stage WHERE item_id=:token AND stage "
                    "LIKE 'PUBLICATION_RANGE_%'"
                ),
                {"token": state.token},
            )
            == 2
        )


@pytest.mark.asyncio
async def test_partial_cancel_recovery_preserves_its_committed_authoritative_flags(db, monkeypatch):
    from juya_admin_api.modules.content_production import commands

    order, _ = await imported_order(db, 2)
    first, second = (await db.service.items(order.work_order_id)).items
    delegate = commands.delegate

    async def lost_after_cancel(*args, **kwargs):
        await delegate(*args, **kwargs)
        raise RuntimeError("outer result lost after local cancellation committed")

    monkeypatch.setattr(commands, "delegate", lost_after_cancel)
    with pytest.raises(RuntimeError):
        await db.service.command(
            order.work_order_id,
            "cancel",
            {
                "expected_version": order.version,
                "scope": "selected",
                "selected_item_ids": [first.item_id],
            },
            db.actor,
            "cancel-one",
        )
    await expire_command(db, order.work_order_id)
    monkeypatch.setattr(commands, "delegate", delegate)
    pending = (await db.service.get(order.work_order_id)).pending_commands
    result = await db.service.resume(order.work_order_id, pending[0].command_id, db.actor)
    assert result.selected_item_ids == [first.item_id] and result.version == order.version + 1
    page = await db.service.items(order.work_order_id)
    assert page.items[0].production_status == "cancelled" and page.items[0].content_id is None
    assert page.items[1].item_id == second.item_id
    assert page.items[1].production_status == "review_required"
