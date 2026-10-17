"""Real minimum XLSX source fields bind server identities and reject repurposed import keys."""

import json
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy import text
from test_v15_import_flow import import_db, validated  # noqa: F401


def source_package(
    external,
    *,
    meaning="河岸",
    source="A bank by the river.",
    meaning_change="",
    expected=None,
    import_key="bank-river",
    entry_type="word",
):
    document = Document()
    document.add_heading("河岸|A bank", level=1)
    document.add_paragraph("A: A bank by the river.")
    document.add_paragraph("B: I can see the bank.")
    document.add_paragraph("甲:河边有一处岸。")
    document.add_paragraph("乙:我能看到河岸。")
    doc = BytesIO()
    document.save(doc)
    content = Workbook()
    content.active.append(["import_content_id", "document_path", "expected_content_version"])
    content.active.append([external, "documents/scene.docx", expected])
    words = Workbook()
    words.active.append(
        [
            "import_content_id",
            "import_entry_id",
            "entry_type",
            "english",
            "meaning_zh",
            "source_sentence",
            "meaning_change",
        ]
    )
    words.active.append([external, import_key, entry_type, "bank", meaning, source, meaning_change])
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for filename, book in (("内容表.xlsx", content), ("词汇表.xlsx", words)):
            data = BytesIO()
            book.save(data)
            archive.writestr(filename, data.getvalue())
        archive.writestr("documents/scene.docx", doc.getvalue())
    return output.getvalue()


async def mapped(db, external):
    async with db.sessions() as session:
        return await db.service.current_mapping(session, "dialogue", external)


async def confirm(db, preview, key):
    acknowledged = list(
        dict.fromkeys(default for item in preview["items"] for default in item["defaults"])
    )
    assert "default_cover" in acknowledged
    return await db.service.confirm(
        preview["batch_id"],
        preview["preview_hash"],
        preview["preview_version"],
        [item["item_id"] for item in preview["items"]],
        acknowledged,
        [],
        db.actor,
        key,
    )


async def saved_entry(db, revision_id):
    async with db.sessions() as session:
        snapshot = await session.scalar(
            text("SELECT content_snapshot FROM scene_revision WHERE public_id=:id"),
            {"id": revision_id},
        )
    snapshot = json.loads(snapshot) if isinstance(snapshot, str) else snapshot
    return snapshot["vocabulary"][0]


@pytest.mark.asyncio
async def test_minimum_source_fields_stable_identity_wording_and_reject_same_key_repurpose(
    import_db,  # noqa: F811 - pytest registers the imported fixture by this name
):
    db = import_db
    external = uuid4().hex
    _created, preview = await validated(db, source_package(external), "first-source")
    result = await confirm(db, preview, "first-confirm")
    assert result["items"][0]["status"] == "SUCCESS", result
    first = await saved_entry(db, result["items"][0]["revision_id"])
    assert first["source_sentence_ids"] == ["sentence-0001"]
    assert first["source_locator"] == {"sentence_id": "sentence-0001"}
    assert first["entry_id"] and first["sense_id"] and first["source_id"]
    original_identity = tuple(first[field] for field in ("entry_id", "sense_id", "source_id"))
    current = await mapped(db, external)
    _wording, wording = await validated(
        db,
        source_package(
            external,
            meaning="河边的岸",
            meaning_change="wording",
            expected=current["current_version"],
        ),
        "wording",
    )
    assert all(
        issue["code"] != "IMPORT_ENTRY_IDENTITY_CHANGE_REQUIRED"
        for item in wording["items"]
        for issue in item["issues"]
    )
    updated = await confirm(db, wording, "wording-confirm")
    assert updated["items"][0]["status"] == "SUCCESS", updated
    worded = await saved_entry(db, updated["items"][0]["revision_id"])
    assert (
        tuple(worded[field] for field in ("entry_id", "sense_id", "source_id")) == original_identity
    )
    for change in (
        {"meaning": "银行", "meaning_change": "new_sense"},
        {"meaning": "河边的岸", "source": "I can see the bank."},
        {"meaning": "河边的岸", "entry_type": "chunk"},
    ):
        current = await mapped(db, external)
        _changed, changed = await validated(
            db, source_package(external, expected=current["current_version"], **change), uuid4().hex
        )
        assert any(
            issue["code"] == "IMPORT_ENTRY_IDENTITY_CHANGE_REQUIRED"
            and issue["severity"] == "ERROR"
            for item in changed["items"]
            for issue in item["issues"]
        ), changed
        rejected = await confirm(db, changed, uuid4().hex)
        assert rejected["items"][0]["status"] == "FAILED"
        assert rejected["items"][0]["error_code"] == "IMPORT_ENTRY_IDENTITY_CHANGE_REQUIRED"
        assert (await mapped(db, external))["draft_revision_id"] == current["draft_revision_id"]
    async with db.sessions() as session:
        mapping = (
            await session.execute(
                text(
                    "SELECT entry_id,sense_id,source_id FROM import_entry_mapping "
                    "WHERE import_entry_id='bank-river' AND content_id=:content"
                ),
                {"content": current["content_id"]},
            )
        ).one()
        assert tuple(mapping) == original_identity


@pytest.mark.asyncio
async def test_generated_semantic_replacement_key_reimports_its_own_identity(import_db):  # noqa: F811
    from juya_admin_api.modules.content.domain import SceneRevision
    from juya_admin_api.modules.content.repository import SQLAlchemyContentRepository
    from juya_admin_api.modules.content.service import ContentService

    db = import_db
    external = uuid4().hex
    _, preview = await validated(db, source_package(external), "generated-first")
    result = await confirm(db, preview, "generated-confirm")
    revision_id = result["items"][0]["revision_id"]
    content = ContentService(SQLAlchemyContentRepository(db.sessions))
    first = await content.get_revision(revision_id)
    old = dict(first.content["vocabulary"][0])
    first.content["vocabulary"][0].update(chinese="银行", meaning_change="new_sense")
    changed = await db.service.production.save_revision(
        SceneRevision(first.id, first.scene_id, None, content=first.content, created_by=db.actor),
        first.version,
    )
    new = changed.content["vocabulary"][0]
    async with db.sessions() as session:
        key = await session.scalar(
            text("SELECT import_entry_id FROM content_entry WHERE public_id=:id"),
            dict(id=new["entry_id"]),
        )
    current = await mapped(db, external)
    _, preview = await validated(
        db,
        source_package(
            external, meaning="银行", import_key=key, expected=current["current_version"]
        ),
        "generated-reimport",
    )
    result = await confirm(db, preview, "generated-reconfirm")
    assert result["items"][0]["status"] == "SUCCESS", result
    replayed = await saved_entry(db, result["items"][0]["revision_id"])
    assert all(
        replayed[field] == new[field] != old[field]
        for field in ("entry_id", "sense_id", "source_id")
    )
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text(
                    "SELECT entry_id FROM import_entry_mapping WHERE "
                    "namespace='juya-content-library' AND content_type='dialogue' AND "
                    "content_id=:content AND import_entry_id='bank-river'"
                ),
                dict(content=first.scene_id),
            )
            == old["entry_id"]
        )
