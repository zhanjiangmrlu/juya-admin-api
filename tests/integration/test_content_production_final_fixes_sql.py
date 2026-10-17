"""Final cross-boundary regressions on private migrated SQL."""

from uuid import uuid4

import pytest
from sqlalchemy import text

from juya_admin_api.modules.content.production_store import encode
from juya_admin_api.modules.content.v15_proofreading import snapshot_hash
from juya_admin_api.shared.errors import AppError
from tests.integration.test_content_work_orders import db as work_db
from tests.integration.test_content_work_orders import manual, reading

db = work_db


async def command(db, order, action, **fields):
    current = await db.service.get(order.work_order_id)
    return await db.service.command(
        order.work_order_id,
        action,
        {"expected_version": current.version, **fields},
        db.actor,
        uuid4().hex,
    )


@pytest.mark.asyncio
async def test_dynamic_and_owned_historical_acknowledgement_contract(db):
    body = reading()
    body.update(category="", copyright_status="")
    order = await manual(db, content=body)
    item = (await db.service.items(order.work_order_id)).items[0]
    source = {"blocks": [{"text": "Original"}]}
    async with db.sessions() as session, session.begin():
        await session.execute(
            text(
                "UPDATE content_work_item SET source_snapshot=:s,source_hash=:h,historical_issues"
                "=:i WHERE public_id=:id"
            ),
            dict(
                s=encode(source),
                h=snapshot_hash(source),
                i=encode(
                    [
                        dict(
                            code="IMPORT_LEXICON_AMBIGUOUS",
                            field="vocabulary",
                            message="核对附表",
                            severity="DRAFT",
                        )
                    ]
                ),
                id=item.item_id,
            ),
        )
    detail = await db.service.item(order.work_order_id, item.item_id)
    dynamic = [i for i in detail.issues if i.code in {"CATEGORY_REQUIRED", "COPYRIGHT_REQUIRED"}]
    assert len(dynamic) == 2
    assert all(
        not i.acknowledgement_eligible and i.location and i.message != "请核对当前内容与原始输入"
        for i in dynamic
    )
    owned = next(i for i in detail.issues if i.code == "IMPORT_LEXICON_AMBIGUOUS")
    assert owned.acknowledgement_eligible and owned.origin_work_item_id == item.item_id
    assert owned.source_hash == snapshot_hash(source)
    result = await db.service.correct(
        order.work_order_id,
        item.item_id,
        dict(
            expected_version=order.version,
            expected_revision_id=item.revision_id,
            expected_edit_version=item.edit_version,
            expected_content_version=item.content_version,
            candidate_hash=item.candidate_hash,
            content=reading(),
            original_compared=True,
            resolved_issue_ids=[owned.id],
        ),
        db.actor,
        "repair",
    )
    assert result.status == "SUCCEEDED"
    assert (await db.service.item(order.work_order_id, item.item_id)).issues == []


async def members(db, count=3):
    orders = [await manual(db, key=uuid4().hex, content=reading(str(n))) for n in range(count)]
    rows = [(await db.service.items(o.work_order_id)).items[0] for o in orders]
    async with db.sessions() as session, session.begin():
        for index, row in enumerate(rows[1:], 2):
            await session.execute(
                text(
                    "UPDATE content_work_item SET work_order_id=:o,position=:p WHERE public_id=:id"
                ),
                dict(o=orders[0].work_order_id, p=index, id=row.item_id),
            )
    return orders[0], rows


@pytest.mark.asyncio
async def test_fresh_sampling_excludes_no_draft_but_keeps_published_content_draft(db):
    order, rows = await members(db)
    old = await command(db, order, "sampling", scope="all")
    for row in rows:
        await db.proof.confirm(
            old.proofreading_set_id,
            row.revision_id,
            row.candidate_hash,
            db.actor,
            {"text_checked": True},
        )
    await db.proof.complete(old.proofreading_set_id, db.actor)
    original = await db.proof.get(old.proofreading_set_id)
    published = []
    for row in rows[:2]:
        published.append(
            await db.readings.command(
                "publish",
                row.content_id,
                db.actor,
                uuid4().hex,
                dict(
                    expected_version=row.content_version,
                    revision_id=row.revision_id,
                    manual_confirmed=True,
                ),
            )
        )
    await db.readings.command(
        "revision",
        rows[1].content_id,
        db.actor,
        uuid4().hex,
        dict(expected_version=published[1]["version"], source_revision_id=rows[1].revision_id),
    )
    rows[1] = await db.service.item(order.work_order_id, rows[1].item_id)
    assert rows[1].published and rows[1].current_draft
    new = await command(db, order, "sampling", scope="all")
    assert new.proofreading_set_id != old.proofreading_set_id
    fresh = await db.proof.get(new.proofreading_set_id)
    assert {c["revision_id"] for c in fresh["candidates"]} == {
        rows[1].revision_id,
        rows[2].revision_id,
    }
    assert (await db.proof.get(old.proofreading_set_id))["candidates"] == original["candidates"]


@pytest.mark.asyncio
@pytest.mark.parametrize("affected_count", [1, 2])
async def test_exact_anomaly_subset_and_replay_leave_unrelated_approval(db, affected_count):
    order, rows = await members(db)
    sampled = await command(db, order, "sampling", scope="all")
    for row in rows:
        await db.proof.confirm(
            sampled.proofreading_set_id,
            row.revision_id,
            row.candidate_hash,
            db.actor,
            {"text_checked": True},
        )
    await db.proof.complete(sampled.proofreading_set_id, db.actor)
    state = await db.proof.get(sampled.proofreading_set_id)
    body = dict(
        expected_version=(await db.service.get(order.work_order_id)).version,
        scope="selected",
        selected_item_ids=[r.item_id for r in rows[:affected_count]],
        operation="fail",
        observed_item_id=rows[0].item_id,
        candidate_hash=rows[0].candidate_hash,
        collection_hash=state["collection_hash"],
        error_code="MANUAL_REVIEW_REQUIRED",
    )
    result = await db.service.command(
        order.work_order_id, "sample-action", body, db.actor, "fail-subset"
    )
    assert (
        await db.service.command(
            order.work_order_id, "sample-action", body, db.actor, "fail-subset"
        )
        == result
    )
    after = await db.proof.get(sampled.proofreading_set_id)
    assert after["status"] == "PARTIAL"
    assert (await db.service.sampling(order.work_order_id)).normal_count == 3 - affected_count
    assert {c["revision_id"] for c in after["candidates"] if c.get("failure")} == {
        r.revision_id for r in rows[:affected_count]
    }
    async with db.sessions() as session:
        confirmed = list(
            (
                await session.execute(
                    text("SELECT revision_id FROM proofreading_confirmation WHERE set_id=:id"),
                    dict(id=sampled.proofreading_set_id),
                )
            ).scalars()
        )
    assert set(confirmed) == {r.revision_id for r in rows[affected_count:]}
    for affected in rows[:affected_count]:
        pending = await db.service.item(order.work_order_id, affected.item_id)
        assert not pending.publishable and pending.production_status == "review_required"
    with pytest.raises(AppError) as incomplete:
        await db.proof.complete(sampled.proofreading_set_id, db.actor)
    assert incomplete.value.code == "PROOFREADING_INCOMPLETE"
    for unaffected in rows[affected_count:]:
        current = await db.service.item(order.work_order_id, unaffected.item_id)
        assert current.publishable
        await db.readings.command(
            "publish",
            current.content_id,
            db.actor,
            uuid4().hex,
            dict(
                expected_version=current.content_version,
                revision_id=current.revision_id,
                manual_confirmed=True,
            ),
        )
        assert (await db.service.item(order.work_order_id, unaffected.item_id)).published
    unchanged = await command(db, order, "sampling", scope="all")
    unchanged_state = await db.proof.get(unchanged.proofreading_set_id)
    assert unchanged_state["normal_count"] == 0 and unchanged_state["sample_ids"] == []
    assert all(c.get("failure") and not c["normal"] for c in unchanged_state["candidates"])
    for affected in rows[:affected_count]:
        current = await db.service.item(order.work_order_id, affected.item_id)
        await db.service.correct(
            order.work_order_id,
            current.item_id,
            dict(
                expected_version=(await db.service.get(order.work_order_id)).version,
                expected_revision_id=current.revision_id,
                expected_edit_version=current.edit_version,
                expected_content_version=current.content_version,
                candidate_hash=current.candidate_hash,
                content=reading("Repaired held draft"),
            ),
            db.actor,
            uuid4().hex,
        )
    repaired = await command(db, order, "sampling", scope="all")
    fresh = await db.proof.get(repaired.proofreading_set_id)
    assert {c["revision_id"] for c in fresh["candidates"]} == {
        i.revision_id for i in rows[:affected_count]
    }
    for candidate in fresh["candidates"]:
        await db.proof.confirm(
            repaired.proofreading_set_id,
            candidate["revision_id"],
            candidate["candidate_hash"],
            db.actor,
            {"text_checked": True},
        )
    await db.proof.complete(repaired.proofreading_set_id, db.actor)
    for affected in rows[:affected_count]:
        assert (await db.service.item(order.work_order_id, affected.item_id)).publishable
    assert (await db.proof.get(sampled.proofreading_set_id))["candidates"] == after["candidates"]


@pytest.mark.asyncio
async def test_anomaly_subset_is_atomic_and_rejects_stale_member(db):
    order, rows = await members(db)
    sampled = await command(db, order, "sampling", scope="all")
    for row in rows:
        await db.proof.confirm(
            sampled.proofreading_set_id,
            row.revision_id,
            row.candidate_hash,
            db.actor,
            {"text_checked": True},
        )
    state = await db.proof.get(sampled.proofreading_set_id)
    candidates = [
        dict(
            content_type="short_reading",
            content_id=i.content_id,
            revision_id=i.revision_id,
            candidate_hash=i.candidate_hash,
        )
        for i in rows[:2]
    ]
    candidates[1]["candidate_hash"] = "f" * 64
    with pytest.raises(AppError) as denied:
        await db.proof.fail(
            sampled.proofreading_set_id,
            rows[0].revision_id,
            "MANUAL_REVIEW_REQUIRED",
            db.actor,
            affected_candidates=candidates,
            collection_digest=state["collection_hash"],
            idempotency_key="atomic-fail",
        )
    assert denied.value.code == "PROOFREADING_HASH_CHANGED"
    assert not any(
        c.get("failure") for c in (await db.proof.get(sampled.proofreading_set_id))["candidates"]
    )
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT COUNT(*) FROM proofreading_confirmation WHERE set_id=:id"),
                dict(id=sampled.proofreading_set_id),
            )
            == 3
        )


@pytest.mark.asyncio
async def test_anomaly_unknown_response_recovers_delegate_receipt(db, monkeypatch):
    from juya_admin_api.modules.content_production import commands
    from tests.integration.test_content_work_order_recovery_sql import expire_command

    order, rows = await members(db)
    sampled = await command(db, order, "sampling", scope="all")
    state = await db.proof.get(sampled.proofreading_set_id)
    body = dict(
        expected_version=(await db.service.get(order.work_order_id)).version,
        scope="selected",
        selected_item_ids=[rows[0].item_id, rows[1].item_id],
        observed_item_id=rows[0].item_id,
        operation="fail",
        candidate_hash=rows[0].candidate_hash,
        collection_hash=state["collection_hash"],
        error_code="MANUAL_REVIEW_REQUIRED",
    )
    original = commands.delegate

    async def dropped(*args, **kwargs):
        await original(*args, **kwargs)
        raise TimeoutError("response lost after atomic delegate commit")

    monkeypatch.setattr(commands, "delegate", dropped)
    with pytest.raises(TimeoutError):
        await db.service.command(order.work_order_id, "sample-action", body, db.actor, "lost-fail")
    monkeypatch.setattr(commands, "delegate", original)
    await expire_command(db, order.work_order_id)
    recovered = await db.service.command(
        order.work_order_id, "sample-action", body, db.actor, "lost-fail"
    )
    assert set(recovered.selected_item_ids) == {rows[0].item_id, rows[1].item_id}
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM audit_event WHERE action='proofreading.fail' "
                    "AND object_public_id=:id"
                ),
                dict(id=sampled.proofreading_set_id),
            )
            == 1
        )


@pytest.mark.asyncio
async def test_correction_fresh_set_listening_reload_and_publish(db):
    from juya_admin_api.shared.errors import AppError
    from tests.integration.test_content_production_closure_sql import prepared_reading

    order, item = await prepared_reading(db)
    old = await command(db, order, "sampling", scope="all")
    body = item.current.content.model_dump(mode="json")
    body["paragraphs_en"] = ["Corrected English."]
    await db.service.correct(
        order.work_order_id,
        item.item_id,
        dict(
            expected_version=(await db.service.get(order.work_order_id)).version,
            expected_revision_id=item.revision_id,
            expected_edit_version=item.edit_version,
            expected_content_version=item.content_version,
            candidate_hash=item.candidate_hash,
            content=body,
        ),
        db.actor,
        "correct-text",
    )
    with pytest.raises(AppError) as stale:
        await db.proof.confirm(
            old.proofreading_set_id,
            item.revision_id,
            item.candidate_hash,
            db.actor,
            {"text_checked": True, "audio_listened": True},
        )
    assert stale.value.code == "PROOFREADING_HASH_CHANGED"
    new = await command(db, order, "sampling", scope="all")
    current = await db.service.item(order.work_order_id, item.item_id)
    await command(
        db,
        order,
        "sample-action",
        scope="selected",
        selected_item_ids=[item.item_id],
        operation="confirm",
        candidate_hash=current.candidate_hash,
        text_checked=True,
        audio_listened=True,
    )
    confirmed = await db.service.item(order.work_order_id, item.item_id)
    assert confirmed.current.content.audio.listening_confirmed
    assert confirmed.current.content.audio.confirmed_asset_id == item.current.content.audio.asset_id
    await command(db, order, "sample-action", operation="complete", allow_abnormal_drafts=True)
    confirmed = await db.service.item(order.work_order_id, item.item_id)
    assert confirmed.publishable
    assert (await db.proof.get(old.proofreading_set_id))["candidates"][0][
        "candidate_hash"
    ] == item.candidate_hash
    assert new.proofreading_set_id != old.proofreading_set_id
    await db.readings.command(
        "publish",
        item.content_id,
        db.actor,
        uuid4().hex,
        dict(
            expected_version=confirmed.content_version,
            revision_id=item.revision_id,
            manual_confirmed=True,
        ),
    )
    published = await db.service.item(order.work_order_id, item.item_id)
    assert published.published and not published.current_draft


@pytest.mark.asyncio
@pytest.mark.parametrize("change", ["paragraphs_en", "asset_id", "version_id"])
async def test_formal_listening_invalidates_after_substantive_correction(db, change):
    from tests.integration.test_content_production_closure_sql import prepared_reading

    order, item = await prepared_reading(db)
    await command(db, order, "sampling", scope="all")
    await command(
        db,
        order,
        "sample-action",
        scope="selected",
        selected_item_ids=[item.item_id],
        operation="confirm",
        candidate_hash=item.candidate_hash,
        text_checked=True,
        audio_listened=True,
    )
    item = await db.service.item(order.work_order_id, item.item_id)
    assert item.current.content.audio.listening_confirmed
    body = item.current.content.model_dump(mode="json")
    if change == "paragraphs_en":
        body[change] = ["Changed text."]
    else:
        body["audio"][change] = uuid4().hex[:26]
    await db.service.correct(
        order.work_order_id,
        item.item_id,
        dict(
            expected_version=(await db.service.get(order.work_order_id)).version,
            expected_revision_id=item.revision_id,
            expected_edit_version=item.edit_version,
            expected_content_version=item.content_version,
            candidate_hash=item.candidate_hash,
            content=body,
        ),
        db.actor,
        uuid4().hex,
    )
    updated = await db.service.item(order.work_order_id, item.item_id)
    assert not updated.confirmed and not updated.current.content.audio.listening_confirmed
    assert updated.current.content.audio.confirmed_asset_id is None


@pytest.mark.asyncio
async def test_docx_ambiguous_entries_can_compare_and_resolve_owned_evidence(db):
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_heading("你好|Hello", level=1)
    doc.add_paragraph("Hello.")
    doc.add_paragraph("你好。")
    doc.add_paragraph("Vocabulary")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "english"
    table.cell(0, 1).text = "meaning_zh"
    table.cell(1, 0).text = "hello"
    table.cell(1, 1).text = "你好"
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "nested original explanation"
    output = BytesIO()
    doc.save(output)
    created = await db.imports.create(
        series_id=db.series,
        content_type="word_card",
        mode="DOCX",
        data=output.getvalue(),
        actor=db.actor,
        key="nested-docx",
    )
    async with db.sessions() as session:
        stage = await session.scalar(
            text("SELECT public_id FROM content_import_stage WHERE item_id=:id"),
            {"id": created["batch_id"]},
        )
    await db.imports.validate(created["batch_id"], stage)
    order = await db.service.from_batch(created["batch_id"], db.actor, "nested-order")
    preview = await db.service.import_preview(order.work_order_id)
    item = (await db.service.items(order.work_order_id)).items[0]
    result = await db.imports.confirm(
        preview.batch_id,
        preview.preview_hash,
        preview.preview_version,
        [item.import_item_id],
        ["default_cover"],
        [],
        db.actor,
        "nested-confirm",
    )
    assert result["items"][0]["status"] == "SUCCESS", result
    item = await db.service.item(order.work_order_id, item.item_id)
    issue = next(i for i in item.issues if i.code == "IMPORT_LEXICON_AMBIGUOUS")
    assert issue.correction_target == "entries" and issue.acknowledgement_eligible
    body = item.current.content.model_dump(mode="json")
    body["vocabulary"] = [
        dict(english="hello", chinese="你好", source_sentence_ids=[body["dialogue"][0]["id"]])
    ]
    await db.service.correct(
        order.work_order_id,
        item.item_id,
        dict(
            expected_version=order.version,
            expected_revision_id=item.revision_id,
            expected_edit_version=item.edit_version,
            expected_content_version=item.content_version,
            candidate_hash=item.candidate_hash,
            content=body,
            original_compared=True,
            resolved_issue_ids=[issue.id],
        ),
        db.actor,
        "repair-entry",
    )
    fresh = await db.service.item(order.work_order_id, item.item_id)
    assert not any(i.code == "IMPORT_LEXICON_AMBIGUOUS" for i in fresh.issues)
    assert fresh.source_snapshot == item.source_snapshot
    assert next(i for i in fresh.historical_issues if i.id == issue.id).resolved
