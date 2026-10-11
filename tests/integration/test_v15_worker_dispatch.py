"""Actual SQL outbox/lease and import stages; broker and cloud are bounded fixtures."""

import os
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
import pytest_asyncio
from docx import Document
from sqlalchemy import text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from juya_admin_api.infrastructure.config import Settings
from juya_admin_api.infrastructure.tasks import v15 as worker
from juya_admin_api.infrastructure.v15_runtime import build_v15_services
from juya_admin_api.modules.content.production_store import decode
from juya_admin_api.modules.media.repository import SQLAlchemyMediaRepository
from juya_admin_api.modules.media.service import MediaService
from juya_admin_api.shared.errors import AppError
from juya_admin_api.shared.ids import new_ulid


class NoCloudOSS:
    async def freeze_bytes(self, *_args):
        raise AssertionError("text-only import must not write a paid or cloud resource")


class NoCloudSafety:
    async def scan_text(self, *_args):
        raise AssertionError("text parsing never calls a cloud provider")


@pytest_asyncio.fixture
async def dispatched_runtime(tmp_path, monkeypatch):
    url = os.getenv("JUYA_TEST_DATABASE_URL")
    if not url:
        pytest.skip("isolated migrated MySQL required")
    engine = create_async_engine(url.replace("mysql+pymysql://", "mysql+asyncmy://"))
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    oss, safety = NoCloudOSS(), NoCloudSafety()
    settings = Settings(
        _env_file=None,
        import_storage_path=str(tmp_path),
        internal_hmac_secret="isolated-worker-credential-key-32bytes",
        content_ai_enabled=False,
        asr_enabled=False,
        tts_enabled=False,
    )
    media = MediaService(oss, SQLAlchemyMediaRepository(sessions), security=safety)
    services = build_v15_services(settings, sessions, oss, media, safety)
    closed = []

    async def close():
        closed.append(True)

    runtime = SimpleNamespace(engine=engine, v15=services, close=close)
    monkeypatch.setattr(worker, "build_runtime", lambda _settings: runtime)
    # No paid pronunciation job may contact a broker during these content queue tests.
    monkeypatch.setattr(worker.pronunciation, "delay", lambda task_id: None)
    actor = new_ulid(datetime.now(UTC))
    series = await services.production.create_series("Worker import", uuid4().hex, None)
    try:
        yield SimpleNamespace(
            sessions=sessions,
            services=services,
            settings=settings,
            actor=actor,
            series=series["id"],
            closed=closed,
        )
    finally:
        await services.client.aclose()
        await engine.dispose()


async def stage_state(db, stage_id):
    async with db.sessions() as session:
        return (
            (
                await session.execute(
                    text(
                        "SELECT status,attempt,fencing_version,lease_until,output,error_code "
                        "FROM content_import_stage WHERE public_id=:id"
                    ),
                    {"id": stage_id},
                )
            )
            .mappings()
            .one()
        )


async def queued_document(db):
    document = Document()
    document.add_heading("邻里|Neighbours", level=1)
    document.add_paragraph("Hello, neighbour.")
    document.add_paragraph("你好\uff0c邻居。")
    stream = BytesIO()
    document.save(stream)
    created = await db.services.imports.create(
        series_id=db.series,
        content_type="short_reading",
        mode="DOCX",
        data=stream.getvalue(),
        actor=db.actor,
        key=uuid4().hex,
    )
    async with db.sessions() as session:
        stage_id = await session.scalar(
            text(
                "SELECT public_id FROM content_import_stage WHERE item_id=:batch "
                "AND stage='VALIDATE'"
            ),
            {"batch": created["batch_id"]},
        )
    return created["batch_id"], stage_id


@pytest.mark.asyncio
async def test_broker_failure_redelivery_queued_expiry_and_actual_import_confirm(
    dispatched_runtime, monkeypatch
):
    db = dispatched_runtime
    batch_id, stage_id = await queued_document(db)

    def unavailable(_stage):
        raise ConnectionError("isolated simulated broker outage")

    monkeypatch.setattr(worker.process_stage, "delay", unavailable)
    await worker.dispatch_async(db.settings)
    assert (await stage_state(db, stage_id))["status"] == "QUEUED"
    async with db.sessions() as session:
        outbox = (
            (
                await session.execute(
                    text(
                        "SELECT status,attempt_count,next_attempt_at FROM admin_outbox "
                        "WHERE JSON_UNQUOTE(JSON_EXTRACT(payload,'$.stage_id'))=:stage"
                    ),
                    {"stage": stage_id},
                )
            )
            .mappings()
            .all()
        )
    assert outbox and all(row["status"] == "PENDING" for row in outbox)
    assert all(row["attempt_count"] == 1 and row["next_attempt_at"] for row in outbox)
    fence_before = (await stage_state(db, stage_id))["fencing_version"]
    async with db.sessions() as session, session.begin():
        await session.execute(
            text(
                "UPDATE content_import_stage SET "
                "lease_until=DATE_SUB(UTC_TIMESTAMP(6),INTERVAL 1 SECOND) "
                "WHERE public_id=:id"
            ),
            {"id": stage_id},
        )
        await session.execute(
            text(
                "UPDATE admin_outbox SET "
                "next_attempt_at=DATE_SUB(UTC_TIMESTAMP(6),INTERVAL 1 SECOND) "
                "WHERE JSON_UNQUOTE(JSON_EXTRACT(payload,'$.stage_id'))=:stage"
            ),
            {"stage": stage_id},
        )
    captured = []
    monkeypatch.setattr(worker.process_stage, "delay", captured.append)
    await worker.dispatch_async(db.settings)
    assert stage_id in captured
    assert (await stage_state(db, stage_id))["fencing_version"] > fence_before
    async with db.sessions() as session:
        states = (
            (
                await session.execute(
                    text(
                        "SELECT status FROM admin_outbox WHERE "
                        "JSON_UNQUOTE(JSON_EXTRACT(payload,'$.stage_id'))=:stage"
                    ),
                    {"stage": stage_id},
                )
            )
            .scalars()
            .all()
        )
    assert states and set(states) == {"PUBLISHED"}

    # Route through the actual worker, actual DOCX parser and SQL compare-and-swap commit.
    await worker.process_stage_async(db.settings, stage_id)
    validation = await stage_state(db, stage_id)
    assert validation["status"] == "SUCCEEDED" and validation["attempt"] == 1
    await worker.process_stage_async(db.settings, stage_id)
    assert (await stage_state(db, stage_id))["attempt"] == 1
    preview = await db.services.imports.preview(batch_id)
    assert len(preview["items"]) == 1
    item = preview["items"][0]
    confirmation = await db.services.imports.queue_confirmation(
        batch_id,
        {
            "preview_hash": preview["preview_hash"],
            "preview_version": preview["preview_version"],
            "selected": [item["item_id"]],
            "acknowledged_defaults": ["default_cover"],
            "confirmed_conflicts": [],
        },
        db.actor,
        uuid4().hex,
    )
    await worker.process_stage_async(db.settings, confirmation["stage_id"])
    completed = await stage_state(db, confirmation["stage_id"])
    assert completed["status"] == "SUCCEEDED", completed
    output = decode(completed["output"])
    assert output["items"][0]["status"] == "SUCCESS", output
    revision = output["items"][0]["revision_id"]
    async with db.sessions() as session:
        reading = await db.services.readings.revision(
            session, output["items"][0]["content_id"], revision
        )
    assert decode(reading["content_snapshot"])["paragraphs_en"] == ["Hello, neighbour."]
    await worker.process_stage_async(db.settings, confirmation["stage_id"])
    assert (await stage_state(db, confirmation["stage_id"]))["attempt"] == 1
    assert len(db.closed) >= 6


@pytest.mark.asyncio
async def test_material_ai_routing_and_expired_worker_fence_cannot_finish(
    dispatched_runtime, monkeypatch
):
    db = dispatched_runtime
    stages = db.services.imports.stages
    routed = []

    async def material(stage_id, owner, now):
        assert owner and now.tzinfo is UTC
        routed.append(("material", stage_id))

    async def candidate(stage_id, owner, now):
        assert owner and now.tzinfo is UTC
        routed.append(("candidate", stage_id))

    monkeypatch.setattr(db.services.materials, "run", material)
    monkeypatch.setattr(db.services.candidates, "run", candidate)
    identifiers = {}
    async with db.sessions() as session, session.begin():
        for kind in ("MATERIAL_VALIDATE", "AI_LEXICON", "ASR_BOUNDARY", "ASR_POLL", "UNSUPPORTED"):
            identifiers[kind] = await stages.ensure_within(
                session, uuid4().hex, kind, uuid4().hex * 2, "routing-fixture"
            )
    for kind, stage_id in identifiers.items():
        await worker.process_stage_async(db.settings, stage_id)
        if kind == "UNSUPPORTED":
            failed = await stage_state(db, stage_id)
            assert failed["status"] == "FAILED"
            assert failed["error_code"] == "IMPORT_STAGE_UNSUPPORTED"
        else:
            assert ("material" if kind.startswith("MATERIAL") else "candidate", stage_id) in routed
    await worker.process_stage_async(db.settings, new_ulid(datetime.now(UTC)))

    fenced_item = uuid4().hex
    async with db.sessions() as session, session.begin():
        fenced = await stages.ensure_within(session, fenced_item, "CONFIRM", "a" * 64, "v1")
    old = await stages.claim(fenced, "old-worker", datetime.now(UTC), 60)
    assert old
    async with db.sessions() as session, session.begin():
        await session.execute(
            text(
                "UPDATE content_import_stage SET "
                "lease_until=DATE_SUB(UTC_TIMESTAMP(6),INTERVAL 1 SECOND) "
                "WHERE public_id=:id"
            ),
            {"id": fenced},
        )
    assert await stages.enqueue(fenced, datetime.now(UTC))
    replacement = await stages.claim(fenced, "replacement-worker", datetime.now(UTC))
    assert replacement and replacement.fencing_version > old.fencing_version
    async with db.sessions() as session, session.begin():
        with pytest.raises(AppError) as lost:
            await db.services.imports.assert_claim_within(session, old, fenced_item)
        assert lost.value.code == "IMPORT_STAGE_LEASE_LOST"
    with pytest.raises(AppError) as stale_output:
        await stages.finish(old, {"obsolete": True}, datetime.now(UTC))
    assert stale_output.value.code == "IMPORT_STAGE_LEASE_LOST"
    await stages.finish(replacement, {"current": True}, datetime.now(UTC))
    assert decode((await stage_state(db, fenced))["output"]) == {"current": True}
