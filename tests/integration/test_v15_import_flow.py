import asyncio
import os
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile

import pytest
import pytest_asyncio
from docx import Document
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from juya_admin_api.modules.content.production_store import ProductionStore
from juya_admin_api.modules.content_import.service import ImportService
from juya_admin_api.modules.short_reading.store import ShortReadingStore
from juya_admin_api.shared.errors import AppError


def import_zip(external, english="Hello", expected=None, *, bad=False, orphan=False):
    document = Document()
    document.add_heading("你好\uff5cHello", level=1)
    document.add_paragraph("A: " + english)
    document.add_paragraph("甲\uff1a你好。")
    doc = BytesIO()
    document.save(doc)
    table = Workbook()
    table.active.append(["import_content_id", "document_path", "expected_content_version"])
    table.active.append([external, "documents/one.docx", expected])
    if bad:
        table.active.append([external + "-bad", "missing.docx", None])
    sheet = BytesIO()
    table.save(sheet)
    files = {"内容表.xlsx": sheet.getvalue(), "documents/one.docx": doc.getvalue()}
    if orphan:
        words = Workbook()
        words.active.append(["import_content_id", "import_entry_id", "entry_type"])
        words.active.append(["unknown", "word", "word"])
        stream = BytesIO()
        words.save(stream)
        files["词汇表.xlsx"] = stream.getvalue()
    package = BytesIO()
    with ZipFile(package, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return package.getvalue()


@pytest_asyncio.fixture
async def import_db(tmp_path):
    url = os.getenv("JUYA_TEST_DATABASE_URL")
    if not url:
        pytest.skip("isolated migrated MySQL required")
    engine = create_async_engine(url.replace("mysql+pymysql://", "mysql+asyncmy://"))
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    production = ProductionStore(sessions)
    series = await production.create_series("import flow", "import-" + uuid4().hex, None)
    actor = uuid4().hex[:26]

    async def bind(data, kind, uploader, role):
        raise AssertionError("These fixtures have no media; no provider may be called")

    service = ImportService(sessions, tmp_path, production, ShortReadingStore(sessions), bind)
    db = SimpleNamespace(sessions=sessions, service=service, series=series["id"], actor=actor)
    try:
        yield db
    finally:
        async with sessions() as session, session.begin():
            ids = (
                (
                    await session.execute(
                        text("SELECT public_id FROM content_import_batch WHERE series_id=:series"),
                        {"series": db.series},
                    )
                )
                .scalars()
                .all()
            )
            for batch_id in ids:
                await session.execute(
                    text(
                        "DELETE FROM admin_outbox WHERE aggregate_public_id=:id OR "
                        "aggregate_public_id IN (SELECT public_id FROM "
                        "content_import_stage WHERE item_id=:id)"
                    ),
                    {"id": batch_id},
                )
                await session.execute(
                    text("DELETE FROM content_import_stage WHERE item_id=:id"), {"id": batch_id}
                )
                await session.execute(
                    text("DELETE FROM content_import_item WHERE batch_id=:id"), {"id": batch_id}
                )
                await session.execute(
                    text("DELETE FROM content_import_batch WHERE public_id=:id"), {"id": batch_id}
                )
            await session.execute(
                text(
                    "DELETE FROM import_content_mapping WHERE content_id IN (SELECT "
                    "s.public_id FROM scene s JOIN content_series cs ON "
                    "cs.id=s.series_id WHERE cs.public_id=:series)"
                ),
                {"series": db.series},
            )
            await session.execute(
                text(
                    "DELETE FROM content_identity WHERE content_id IN (SELECT "
                    "s.public_id FROM scene s JOIN content_series cs ON "
                    "cs.id=s.series_id WHERE cs.public_id=:series)"
                ),
                {"series": db.series},
            )
            await session.execute(
                text(
                    "DELETE s FROM scene s JOIN content_series cs ON "
                    "cs.id=s.series_id WHERE cs.public_id=:series"
                ),
                {"series": db.series},
            )
            await session.execute(
                text("DELETE FROM content_series WHERE public_id=:series"), {"series": db.series}
            )
            await session.execute(
                text("DELETE FROM idempotency_record WHERE actor_id=:actor"), {"actor": actor}
            )
            await session.execute(
                text("DELETE FROM audit_event WHERE actor_public_id=:actor"), {"actor": actor}
            )
        await engine.dispose()


async def validated(db, data, key):
    created = await db.service.create(
        series_id=db.series, content_type="dialogue", mode="ZIP", data=data, actor=db.actor, key=key
    )
    async with db.sessions() as session:
        stage = await session.scalar(
            text(
                "SELECT public_id FROM content_import_stage WHERE item_id=:id AND stage='VALIDATE'"
            ),
            {"id": created["batch_id"]},
        )
    await db.service.validate(created["batch_id"], stage)
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_stage WHERE public_id=:id"), {"id": stage}
            )
            == "SUCCEEDED"
        )
    return created, await db.service.preview(created["batch_id"])


async def confirm(db, preview, key, *, selected=None, acknowledge=True, claim=None):
    return await db.service.confirm(
        preview["batch_id"],
        preview["preview_hash"],
        preview["preview_version"],
        selected or [item["item_id"] for item in preview["items"]],
        ["default_cover"] if acknowledge else [],
        [],
        db.actor,
        key,
        claim,
    )


@pytest.mark.asyncio
async def test_sql_create_validate_preview_ack_partial_skip_update_and_target_cas(import_db):
    db = import_db
    external = uuid4().hex
    data = import_zip(external, bad=True)
    created, preview = await validated(db, data, "create")
    assert created == await db.service.create(
        series_id=db.series,
        content_type="dialogue",
        mode="ZIP",
        data=data,
        actor=db.actor,
        key="create",
    )
    with pytest.raises(AppError):
        await db.service.create(
            series_id=db.series,
            content_type="dialogue",
            mode="ZIP",
            data=import_zip(external, "Different"),
            actor=db.actor,
            key="create",
        )
    assert len(preview["items"]) == 2
    failed = await confirm(db, preview, "unacknowledged", acknowledge=False)
    assert failed["status"] == "PARTIAL"
    assert failed["items"][0]["error_code"] == "IMPORT_DEFAULT_CONFIRMATION_REQUIRED"
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT COUNT(*) FROM import_content_mapping WHERE import_content_id=:id"),
                {"id": external},
            )
            == 0
        )
    partial = await confirm(db, preview, "partial")
    assert [item["status"] for item in partial["items"]] == ["SUCCESS", "FAILED"]
    assert partial == await confirm(db, preview, "partial")
    identity = partial["items"][0]["content_id"]
    old_revision = partial["items"][0]["revision_id"]
    # Reimport the exact same content identity and bytes: SKIP does not create a revision.
    _skipped, same = await validated(db, data, "same-package")
    assert same["items"][0]["operation"] == "SKIP"
    skipped = await confirm(db, same, "skip", selected=[same["items"][0]["item_id"]])
    assert skipped["items"][0]["content_id"] == identity
    assert skipped["items"][0]["revision_id"] == old_revision
    current_version = same["items"][0]["mapping"]["current_version"]
    _changed, changed = await validated(
        db, import_zip(external, "Changed English", current_version), "changed"
    )
    assert changed["items"][0]["operation"] == "UPDATE_DRAFT"
    updated = await confirm(db, changed, "update")
    assert updated["items"][0]["content_id"] == identity
    assert updated["items"][0]["revision_id"] != old_revision
    # The current mapping is serialized without datetime fields and locks the exact target.
    _cas, cas = await validated(db, import_zip(external, "Newer English", 1), "cas")
    async with db.sessions() as session, session.begin():
        await session.execute(
            text("UPDATE scene_revision SET edit_version=edit_version+1 WHERE public_id=:id"),
            {"id": updated["items"][0]["revision_id"]},
        )
    conflict = await confirm(db, cas, "cas-confirm")
    assert conflict["items"][0]["error_code"] == "IMPORT_TARGET_CHANGED"
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT input_hash FROM import_content_mapping WHERE import_content_id=:id"),
                {"id": external},
            )
            == changed["items"][0]["input_hash"]
        )


@pytest.mark.asyncio
async def test_sql_orphan_rows_are_visible_and_block_all_content_submission(import_db):
    db = import_db
    _created, preview = await validated(db, import_zip(uuid4().hex, orphan=True), "orphan")
    assert preview["manifest"]["orphan_entry_rows"] == [2]
    assert preview["batch_issues"][0]["code"] == "IMPORT_ORPHAN_ENTRY"
    with pytest.raises(AppError) as error:
        await confirm(db, preview, "orphan-confirm")
    assert error.value.code == "IMPORT_BATCH_BLOCKED"
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM scene s JOIN content_series cs ON "
                    "cs.id=s.series_id WHERE cs.public_id=:series"
                ),
                {"series": db.series},
            )
            == 0
        )


@pytest.mark.asyncio
async def test_sql_cancel_during_preparation_fences_old_worker_and_keeps_items_unmodified(
    import_db,
):
    db = import_db
    _created, preview = await validated(db, import_zip(uuid4().hex), "cancelled")
    payload = {
        "preview_hash": preview["preview_hash"],
        "preview_version": preview["preview_version"],
        "selected": [preview["items"][0]["item_id"]],
        "acknowledged_defaults": ["default_cover"],
        "confirmed_conflicts": [],
    }
    queued = await db.service.queue_confirmation(preview["batch_id"], payload, db.actor, "queued")
    assert queued == await db.service.queue_confirmation(
        preview["batch_id"], payload, db.actor, "queued"
    )
    claim = await db.service.stages.claim(queued["stage_id"], "worker", datetime.now(UTC))
    assert claim
    original_prepare = db.service.prepare

    async def cancel_during_prepare(*args):
        await db.service.cancel(preview["batch_id"], db.actor)
        return await original_prepare(*args)

    with (
        patch.object(db.service, "prepare", side_effect=cancel_during_prepare),
        pytest.raises(AppError) as error,
    ):
        await confirm(db, preview, "old-worker", claim=claim)
    assert error.value.code in {"IMPORT_STAGE_LEASE_LOST", "IMPORT_PREVIEW_CHANGED"}
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_batch WHERE public_id=:id"),
                {"id": preview["batch_id"]},
            )
            == "CANCELLED"
        )
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_item WHERE public_id=:id"),
                {"id": preview["items"][0]["item_id"]},
            )
            == "PENDING"
        )
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM idempotency_record WHERE "
                    "scope='v15.import.confirm' AND actor_id=:actor AND "
                    "idempotency_key='old-worker'"
                ),
                {"actor": db.actor},
            )
            == 0
        )
    assert await db.service.stages.claim(queued["stage_id"], "late", datetime.now(UTC)) is None


@pytest.mark.asyncio
async def test_sql_expired_validator_cannot_commit_items_or_preview(import_db):
    db = import_db
    created = await db.service.create(
        series_id=db.series,
        content_type="dialogue",
        mode="ZIP",
        data=import_zip(uuid4().hex),
        actor=db.actor,
        key="stale-validation",
    )
    async with db.sessions() as session:
        stage = await session.scalar(
            text(
                "SELECT public_id FROM content_import_stage WHERE item_id=:id AND stage='VALIDATE'"
            ),
            {"id": created["batch_id"]},
        )
    reached = asyncio.Event()
    released = asyncio.Event()
    original = db.service.original

    async def paused(batch):
        data = await original(batch)
        reached.set()
        await released.wait()
        return data

    with patch.object(db.service, "original", side_effect=paused):
        task = asyncio.create_task(db.service.validate(created["batch_id"], stage))
        await asyncio.wait_for(reached.wait(), 10)
        async with db.sessions() as session, session.begin():
            await session.execute(
                text(
                    "UPDATE content_import_stage SET "
                    "lease_until=UTC_TIMESTAMP(6)-INTERVAL 1 SECOND WHERE public_id=:id"
                ),
                {"id": stage},
            )
        replacement = await db.service.stages.claim(stage, "replacement", datetime.now(UTC))
        assert replacement
        released.set()
        with pytest.raises(AppError):
            await task
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT status FROM content_import_batch WHERE public_id=:id"),
                {"id": created["batch_id"]},
            )
            == "PENDING"
        )
        assert (
            await session.scalar(
                text("SELECT COUNT(*) FROM content_import_item WHERE batch_id=:id"),
                {"id": created["batch_id"]},
            )
            == 0
        )
        assert (
            await session.scalar(
                text("SELECT lease_owner FROM content_import_stage WHERE public_id=:id"),
                {"id": stage},
            )
            == "replacement"
        )


@pytest.mark.asyncio
async def test_sql_explicit_new_key_retries_stage_and_does_not_prepare_success_again(import_db):
    db = import_db
    _created, preview = await validated(db, import_zip(uuid4().hex), "retry-create")
    payload = {
        "preview_hash": preview["preview_hash"],
        "preview_version": preview["preview_version"],
        "selected": [preview["items"][0]["item_id"]],
        "acknowledged_defaults": ["default_cover"],
        "confirmed_conflicts": [],
    }
    first = await db.service.queue_confirmation(
        preview["batch_id"], payload, db.actor, "first-queue"
    )
    first_claim = await db.service.stages.claim(
        first["stage_id"], "first-worker", datetime.now(UTC)
    )
    assert first_claim
    with patch.object(db.service, "prepare", side_effect=ValueError("temporary binding failure")):
        failed = await confirm(db, preview, "first-execute", claim=first_claim)
    assert failed["items"][0]["status"] == "FAILED"
    await db.service.stages.finish(first_claim, failed, datetime.now(UTC))
    second = await db.service.queue_confirmation(
        preview["batch_id"], payload, db.actor, "second-queue"
    )
    assert second["stage_id"] != first["stage_id"]
    second_claim = await db.service.stages.claim(
        second["stage_id"], "retry-worker", datetime.now(UTC)
    )
    assert second_claim
    result = await confirm(db, preview, "retry-execute", claim=second_claim)
    assert result["items"][0]["status"] == "SUCCESS"
    await db.service.stages.finish(second_claim, result, datetime.now(UTC))
    third = await db.service.queue_confirmation(
        preview["batch_id"], payload, db.actor, "third-queue"
    )
    assert third["stage_id"] not in {first["stage_id"], second["stage_id"]}
    third_claim = await db.service.stages.claim(
        third["stage_id"], "replay-worker", datetime.now(UTC)
    )
    with patch.object(
        db.service, "prepare", side_effect=AssertionError("success must not rebind")
    ) as prepare:
        replayed = await confirm(db, preview, "third-execute", claim=third_claim)
    assert replayed["items"] == result["items"]
    prepare.assert_not_called()
