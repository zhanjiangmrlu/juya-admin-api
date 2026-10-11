"""E08: actual 1000-row ZIP, durable SQL drafts, single-item rollback and recovery.

Only one deliberate save failure is injected. SQL stores and document/ZIP parsing are real.
Media binding must never run because this package contains no media.
"""

import json
import math
import os
import platform
import tracemalloc
from io import BytesIO
from time import perf_counter
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from sqlalchemy import text
from test_v15_import_flow import import_db as import_db

from juya_admin_api.modules.content_import.parser import read_container
from juya_admin_api.modules.content_import.service import ImportService
from juya_admin_api.shared.errors import AppError


def thousand_zip(prefix: str) -> bytes:
    document = Document()
    document.add_heading("你好\uff5cHello", level=1)
    document.add_paragraph("A: Hello.")
    document.add_paragraph("甲\uff1a你好。")
    document_bytes = BytesIO()
    document.save(document_bytes)
    workbook = Workbook()
    workbook.active.append(["import_content_id", "document_path"])
    for index in range(1000):
        workbook.active.append([f"{prefix}-{index:04d}", "documents/shared.docx"])
    workbook_bytes = BytesIO()
    workbook.save(workbook_bytes)
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("内容表.xlsx", workbook_bytes.getvalue())
        archive.writestr("documents/shared.docx", document_bytes.getvalue())
    return output.getvalue()


async def scene_versions(db):
    async with db.sessions() as session:
        rows = (
            (
                await session.execute(
                    text(
                        "SELECT s.public_id,r.public_id AS revision_id,r.version_no,r.edit_version "
                        "FROM scene s JOIN content_series cs ON cs.id=s.series_id "
                        "JOIN scene_revision r ON r.scene_id=s.id WHERE cs.public_id=:series"
                    ),
                    {"series": db.series},
                )
            )
            .mappings()
            .all()
        )
    return {
        row["public_id"]: (row["revision_id"], row["version_no"], row["edit_version"])
        for row in rows
    }, len(rows)


@pytest.mark.asyncio
async def test_thousand_independent_docx_ids_partial_retry_and_io_free_replay(import_db):
    db = import_db
    collect_memory = os.getenv("JUYA_THOUSAND_METRICS") == "true"
    if collect_memory:
        tracemalloc.start()
    elapsed = {}
    save_ms = []
    prefix = uuid4().hex
    data = thousand_zip(prefix)
    with ZipFile(BytesIO(data)) as archive:
        unpacked_bytes = sum(entry.file_size for entry in archive.infolist())
        member_compressed_bytes = sum(entry.compress_size for entry in archive.infolist())
    counts = {"validation_package_extract": 0, "confirmation_package_extract": 0}

    def validate_container(payload, *args, **kwargs):
        if payload == data:
            counts["validation_package_extract"] += 1
        return read_container(payload, *args, **kwargs)

    def confirm_container(payload, *args, **kwargs):
        if payload == data:
            counts["confirmation_package_extract"] += 1
            assert kwargs.get("spool") is True
        return read_container(payload, *args, **kwargs)

    with (
        patch(
            "juya_admin_api.modules.content_import.parser.read_container",
            side_effect=validate_container,
        ),
        patch(
            "juya_admin_api.modules.content_import.service.read_container",
            side_effect=confirm_container,
        ),
    ):
        created = await db.service.create(
            series_id=db.series,
            content_type="dialogue",
            mode="ZIP",
            data=data,
            actor=db.actor,
            key="thousand-create",
        )
        batch_id = created["batch_id"]
        async with db.sessions() as session:
            stage_id = await session.scalar(
                text(
                    "SELECT public_id FROM content_import_stage WHERE item_id=:id AND "
                    "stage='VALIDATE'"
                ),
                {"id": batch_id},
            )
        assert stage_id is not None
        started = perf_counter()
        await db.service.validate(batch_id, stage_id)
        elapsed["validate_seconds"] = perf_counter() - started
        assert counts["validation_package_extract"] == 1
        # A fresh service resumes the same logical batch; validation redelivery does no new IO.
        resumed = ImportService(
            db.sessions,
            db.service.storage,
            db.service.production,
            db.service.readings,
            db.service.bind_media,
        )
        await resumed.validate(batch_id, stage_id)
        assert counts["validation_package_extract"] == 1
        started = perf_counter()
        preview = await resumed.preview(batch_id)
        elapsed["preview_seconds"] = perf_counter() - started
        assert len(preview["items"]) == 1000
        assert len({item["item_id"] for item in preview["items"]}) == 1000
        assert len({item["import_content_id"] for item in preview["items"]}) == 1000
        assert all(item["operation"] == "CREATE" for item in preview["items"])
        assert not preview["batch_issues"]
        selected = [item["item_id"] for item in preview["items"]]
        acknowledged = sorted({value for item in preview["items"] for value in item["defaults"]})
        original_save = resumed.save_item_within
        failed_external = f"{prefix}-0499"
        injected = 0

        async def save_with_one_failure(session, batch, item, mapping, actor):
            nonlocal injected
            started = perf_counter()
            try:
                result = await original_save(session, batch, item, mapping, actor)
            finally:
                save_ms.append((perf_counter() - started) * 1000)
            if item["import_content_id"] == failed_external and injected == 0:
                injected += 1
                # Fail after actual inserts: the item's transaction must remove its new rows.
                raise AppError("THOUSAND_ITEM_FAILURE", "受控单项事务故障", 409)
            return result

        async def issue(active_preview, key, item_ids):
            return await resumed.confirm(
                batch_id,
                active_preview["preview_hash"],
                active_preview["preview_version"],
                item_ids,
                acknowledged,
                [],
                db.actor,
                key,
            )

        with patch.object(resumed, "save_item_within", side_effect=save_with_one_failure):
            started = perf_counter()
            first = await issue(preview, "thousand-confirm", selected)
            elapsed["confirm_1000_seconds"] = perf_counter() - started
        assert injected == 1
        assert first["status"] == "PARTIAL"
        assert len(first["items"]) == 1000
        assert sum(item["status"] == "SUCCESS" for item in first["items"]) == 999
        failed = [item for item in first["items"] if item["status"] == "FAILED"]
        assert len(failed) == 1 and failed[0]["error_code"] == "THOUSAND_ITEM_FAILURE"
        assert counts["confirmation_package_extract"] == 1  # One spool for all 1000 items.
        before, revision_count = await scene_versions(db)
        assert len(before) == revision_count == 999
        async with db.sessions() as session:
            mappings = (
                await session.execute(
                    text(
                        "SELECT import_content_id,content_id FROM import_content_mapping W"
                        "HERE import_content_id LIKE :prefix"
                    ),
                    {"prefix": prefix + "-%"},
                )
            ).all()
        assert len(mappings) == 999
        assert failed_external not in {row.import_content_id for row in mappings}
        assert len({row.content_id for row in mappings}) == 999
        assert await issue(preview, "thousand-confirm", selected) == first
        assert counts["confirmation_package_extract"] == 1
        retry_preview = await resumed.preview(batch_id)
        assert sum(item["operation"] == "SKIP" for item in retry_preview["items"]) == 999
        assert sum(item["operation"] == "CREATE" for item in retry_preview["items"]) == 1
        retry_selected = [failed[0]["item_id"]]
        with patch.object(resumed, "save_item_within", wraps=original_save) as retry_save:
            started = perf_counter()
            retry = await issue(retry_preview, "thousand-retry-one", retry_selected)
            elapsed["recover_one_failed_seconds"] = perf_counter() - started
            assert retry_save.await_count == 1
        assert retry["status"] == "CONFIRMED"
        assert len(retry["items"]) == 1 and retry["items"][0]["status"] == "SUCCESS"
        assert counts["confirmation_package_extract"] == 2
        after, revision_count = await scene_versions(db)
        assert len(after) == revision_count == 1000
        assert all(after[content_id] == version for content_id, version in before.items())
        async with db.sessions() as session:
            final = (
                await session.execute(
                    text(
                        "SELECT import_content_id,content_id FROM import_content_mapping W"
                        "HERE import_content_id LIKE :prefix"
                    ),
                    {"prefix": prefix + "-%"},
                )
            ).all()
            success_count = await session.scalar(
                text(
                    "SELECT COUNT(*) FROM content_import_item WHERE batch_id=:id AND s"
                    "tatus='SUCCESS'"
                ),
                {"id": batch_id},
            )
            batch = await resumed.batch(session, batch_id)
        assert (
            len(final)
            == len({row.import_content_id for row in final})
            == len({row.content_id for row in final})
            == 1000
        )
        assert success_count == 1000
        # Delete the package: committed receipts must replay before any archive/media access.
        resumed.path(batch["asset_id"]).unlink()
        with (
            patch.object(
                resumed, "original", side_effect=AssertionError("Replay reopened archive")
            ),
            patch.object(
                resumed, "bind_media", side_effect=AssertionError("Replay repeated media IO")
            ),
        ):
            assert await issue(preview, "thousand-confirm", selected) == first
            assert await issue(retry_preview, "thousand-retry-one", retry_selected) == retry
        assert counts == {"validation_package_extract": 1, "confirmation_package_extract": 2}
        metadata_ms = []
        async with db.sessions() as session:
            for _ in range(20):
                started = perf_counter()
                metadata = (
                    await session.execute(
                        text(
                            "SELECT status,preview_version FROM content_import_batch "
                            "WHERE public_id=:id"
                        ),
                        {"id": batch_id},
                    )
                ).one()
                assert metadata.status == "CONFIRMED"
                metadata_ms.append((perf_counter() - started) * 1000)
            actual_sentences = await session.scalar(
                text(
                    "SELECT SUM(JSON_LENGTH(r.content_snapshot,'$.dialogue')) "
                    "FROM scene_revision r "
                    "JOIN scene s ON s.id=r.scene_id JOIN content_series cs ON cs.id=s.series_id "
                    "WHERE cs.public_id=:series"
                ),
                {"series": db.series},
            )
            assert actual_sentences == 1000
        peak = tracemalloc.get_traced_memory()[1] if collect_memory else None
        if collect_memory:
            tracemalloc.stop()
        print(
            "V15_THOUSAND_METRICS="
            + json.dumps(
                {
                    "os": platform.system() + " " + platform.release(),
                    "cpu": platform.processor(),
                    "logical_cpus": os.cpu_count(),
                    "concurrency": 1,
                    "zip_archive_bytes": len(data),
                    "zip_member_compressed_bytes": member_compressed_bytes,
                    "zip_uncompressed_bytes": unpacked_bytes,
                    "articles": 1000,
                    "logical_sentences": int(actual_sentences),
                    "media_files": 0,
                    "media_duration_ms": 0,
                    "known_fixture_cloud_calls": 0,
                    "elapsed": elapsed,
                    "python_tracemalloc_peak_bytes": peak,
                    "process_working_set_peak_bytes": None,
                    "metadata_sql_latency_ms": {"n": 20, "p95": sorted(metadata_ms)[18]},
                    "item_save_sql_latency_ms": {
                        "n": len(save_ms),
                        "p95": sorted(save_ms)[math.ceil(len(save_ms) * 0.95) - 1],
                    },
                    "http_latency_p95": None,
                    "whole_task_p95": None,
                    "extraction_counts": counts,
                },
                ensure_ascii=False,
                sort_keys=True,
            )
        )
