"""Missing material replay counterexamples against runner-owned real MySQL."""

import hashlib
import json
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from PIL.PngImagePlugin import PngInfo
from sqlalchemy import text
from test_v15_import_entry_identity import confirm, mapped, saved_entry, source_package
from test_v15_import_flow import import_db, validated  # noqa: F401


def replace_document(package, *, body=None, evidence=False):
    result = BytesIO()
    with ZipFile(BytesIO(package)) as original, ZipFile(result, "w") as changed:
        for name in original.namelist():
            data = original.read(name)
            if name.endswith("scene.docx"):
                document = Document(BytesIO(data))
                if body:
                    document.paragraphs[2].text = body
                if evidence:
                    document.core_properties.author = "Evidence-only correction"
                    document.add_paragraph("")
                data = BytesIO()
                document.save(data)
                data = data.getvalue()
            changed.writestr(name, data)
    return result.getvalue()


async def item_hash(db, batch_id):
    async with db.sessions() as session:
        return await session.scalar(
            text(
                "SELECT JSON_UNQUOTE(JSON_EXTRACT(candidate,'$.semantic_hash')) "
                "FROM content_import_item WHERE batch_id=:id"
            ),
            dict(id=batch_id),
        )


@pytest.mark.asyncio
async def test_trusted_zip_evidence_only_skips_but_body_change_revises_candidate(import_db):  # noqa: F811
    db = import_db
    external = uuid4().hex
    data = source_package(external)
    created, preview = await validated(db, data, "semantic-first")
    first = await confirm(db, preview, "semantic-save")
    assert first["items"][0]["status"] == "SUCCESS", first
    revision = first["items"][0]["revision_id"]
    first_hash = await item_hash(db, created["batch_id"])
    assert first_hash is not None
    async with db.sessions() as session, session.begin():
        old = await session.scalar(
            text("SELECT content_snapshot FROM scene_revision WHERE public_id=:id"),
            dict(id=revision),
        )
        raw_hash = await session.scalar(
            text("SELECT input_hash FROM import_content_mapping WHERE import_content_id=:id"),
            dict(id=external),
        )
        paid_before = await session.scalar(text("SELECT COUNT(*) FROM content_cloud_attempt"))
        # Actual old success envelope: preserve literal mapping and remove only the new field.
        await session.execute(
            text(
                "UPDATE content_import_item SET candidate=JSON_REMOVE(candidate,'$.semantic_hash') "
                "WHERE batch_id=:id"
            ),
            dict(id=created["batch_id"]),
        )
    _, literal = await validated(db, data, "semantic-old-literal")
    assert literal["items"][0]["operation"] == "SKIP"
    evidence, preview = await validated(
        db, replace_document(data, evidence=True), "semantic-evidence"
    )
    assert evidence["input_hash"] != created["input_hash"]
    assert await item_hash(db, evidence["batch_id"]) == first_hash
    assert preview["items"][0]["operation"] == "SKIP"
    skipped = await confirm(db, preview, "semantic-evidence-save")
    assert skipped["items"][0]["revision_id"] == revision
    async with db.sessions() as session:
        assert (
            await session.scalar(
                text("SELECT input_hash FROM import_content_mapping WHERE import_content_id=:id"),
                dict(id=external),
            )
            == raw_hash
        )
        assert (
            await session.scalar(text("SELECT COUNT(*) FROM content_cloud_attempt")) == paid_before
        )
    current = await mapped(db, external)
    changed, preview = await validated(
        db,
        replace_document(
            source_package(external, expected=current["current_version"]),
            body="B: We can see the bank.",
        ),
        "semantic-body",
    )
    assert await item_hash(db, changed["batch_id"]) != first_hash
    assert preview["items"][0]["operation"] == "UPDATE_DRAFT"
    revised = await confirm(db, preview, "semantic-body-save")
    assert revised["items"][0]["status"] == "SUCCESS", revised
    new_revision = revised["items"][0]["revision_id"]
    assert new_revision != revision
    async with db.sessions() as session:
        latest = await session.scalar(
            text("SELECT content_snapshot FROM scene_revision WHERE public_id=:id"),
            dict(id=new_revision),
        )
        assert (
            await session.scalar(
                text("SELECT content_snapshot FROM scene_revision WHERE public_id=:id"),
                dict(id=revision),
            )
            == old
        )
        assert (
            await session.scalar(
                text(
                    "SELECT COUNT(*) FROM scene_revision WHERE scene_id="
                    "(SELECT id FROM scene WHERE public_id=:id)"
                ),
                dict(id=first["items"][0]["content_id"]),
            )
            == 2
        )
    latest = json.loads(latest) if isinstance(latest, str) else latest
    assert latest["dialogue"][1]["english"] == "We can see the bank."


@pytest.mark.asyncio
@pytest.mark.parametrize("untrusted", ["namespace", "type", "key", "content", "hash", "missing"])
async def test_evidence_replay_requires_exact_successful_scoped_mapping(import_db, untrusted):  # noqa: F811
    db = import_db
    external = uuid4().hex
    data = source_package(external)
    created, preview = await validated(db, data, "scope-first")
    first = await confirm(db, preview, "scope-save")
    assert first["items"][0]["status"] == "SUCCESS", first
    async with db.sessions() as session, session.begin():
        if untrusted in {"namespace", "type"}:
            column, value = (
                ("namespace", "other-library")
                if untrusted == "namespace"
                else ("content_type", "word_card")
            )
            await session.execute(
                text(f"UPDATE content_import_batch SET {column}=:value WHERE public_id=:id"),
                dict(value=value, id=created["batch_id"]),
            )
        else:
            column, value = {
                "key": ("import_content_id", "unrelated-key"),
                "content": ("content_id", "unrelated-content"),
                "hash": ("input_hash", "b" * 64),
                "missing": ("candidate", "{}"),
            }[untrusted]
            await session.execute(
                text(f"UPDATE content_import_item SET {column}=:value WHERE batch_id=:id"),
                dict(value=value, id=created["batch_id"]),
            )
    _, changed = await validated(db, replace_document(data, evidence=True), "scope-new")
    assert changed["items"][0]["operation"] == "UPDATE_DRAFT"
    assert (await mapped(db, external))["draft_revision_id"] == first["items"][0]["revision_id"]


def package_cover(package, color, *, identity=""):
    output, image = BytesIO(), BytesIO()
    metadata = PngInfo()
    metadata.add_text("isolated_fixture", identity)
    Image.new("RGB", (16, 16), color).save(image, format="PNG", pnginfo=metadata)
    with ZipFile(BytesIO(package)) as old, ZipFile(output, "w") as new:
        for name in old.namelist():
            data = old.read(name)
            if name == "内容表.xlsx":
                workbook = load_workbook(BytesIO(data))
                workbook.active.cell(1, 4, "cover_path")
                workbook.active.cell(2, 4, "cover.png")
                sheet = BytesIO()
                workbook.save(sheet)
                data = sheet.getvalue()
            new.writestr(name, data)
        new.writestr("cover.png", image.getvalue())
    return output.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("change", "cached_semantic"),
    [
        ("replacement", True),
        ("partial_evidence_removal", False),
        ("partial_evidence_removal", True),
    ],
)
async def test_actual_media_bytes_change_cannot_semantically_skip(
    import_db,  # noqa: F811
    change,
    cached_semantic,
):
    db = import_db
    external = uuid4().hex
    bound_bytes = []

    async def bind_image(data, kind, actor, role):
        assert (kind, actor, role) == ("images", db.actor, "COVER_IMAGE")
        Image.open(BytesIO(data)).verify()
        identity = uuid4().hex[:26]
        async with db.sessions() as session, session.begin():
            await session.execute(
                text(
                    "INSERT INTO media_asset(public_id,object_key,asset_type,content_type,"
                    "size_bytes,sha256,status,security_status,created_by,created_at) "
                    "VALUES(:id,:object,'images','image/png',:size,:sha,'CONFIRMED','PASSED',"
                    ":actor,UTC_TIMESTAMP(6))"
                ),
                dict(
                    id=identity,
                    object="sealed/test/" + identity,
                    size=len(data),
                    sha=hashlib.sha256(data).hexdigest(),
                    actor=actor,
                ),
            )
        bound_bytes.append(data)
        return dict(asset_id=identity)

    db.service.bind_media = bind_image

    async def save_media(preview, key):
        return await db.service.confirm(
            preview["batch_id"],
            preview["preview_hash"],
            preview["preview_version"],
            [item["item_id"] for item in preview["items"]],
            list({value for item in preview["items"] for value in item["defaults"]}),
            [],
            db.actor,
            key,
        )

    data = source_package(external)
    created, preview = await validated(
        db, package_cover(data, "red", identity=external), "media-first"
    )
    first = await save_media(preview, "media-save")
    assert first["items"][0]["status"] == "SUCCESS", first
    original_semantic = await item_hash(db, created["batch_id"])
    current = await mapped(db, external)
    package = source_package(external, expected=current["current_version"])
    if change == "replacement":
        package = package_cover(package, "blue", identity=external)
    else:
        # Real historical row retains its trusted indexed hash/scope/result, but loses one resource.
        paths = ["$.resources.cover_path"]
        if not cached_semantic:
            paths.append("$.semantic_hash")
        async with db.sessions() as session, session.begin():
            await session.execute(
                text(
                    "UPDATE content_import_item SET candidate=JSON_REMOVE(candidate,"
                    + ",".join(f":path{index}" for index in range(len(paths)))
                    + ") WHERE batch_id=:id"
                ),
                {"id": created["batch_id"]}
                | {f"path{index}": path for index, path in enumerate(paths)},
            )
    changed, preview = await validated(db, package, "media-changed")
    if change == "partial_evidence_removal" and cached_semantic:
        # Even a cached semantic value forged to match the new request cannot prove old materials.
        async with db.sessions() as session, session.begin():
            await session.execute(
                text(
                    "UPDATE content_import_item SET candidate=JSON_SET(candidate,"
                    "'$.semantic_hash',:semantic) WHERE batch_id=:id"
                ),
                dict(id=created["batch_id"], semantic=await item_hash(db, changed["batch_id"])),
            )
        preview = await db.service.preview(changed["batch_id"])
    assert await item_hash(db, changed["batch_id"]) != original_semantic
    assert preview["items"][0]["operation"] == "UPDATE_DRAFT"
    second = await save_media(preview, "media-new-save")
    assert second["items"][0]["status"] == "SUCCESS", second
    if change == "replacement":
        assert len(bound_bytes) == 2 and bound_bytes[0] != bound_bytes[1]
    else:
        assert len(bound_bytes) == 1
        async with db.sessions() as session:
            snapshots = (
                await session.execute(
                    text(
                        "SELECT public_id,JSON_UNQUOTE("
                        "JSON_EXTRACT(content_snapshot,'$.cover_asset_id')) "
                        "AS cover FROM scene_revision WHERE public_id IN (:old,:new)"
                    ),
                    dict(
                        old=first["items"][0]["revision_id"], new=second["items"][0]["revision_id"]
                    ),
                )
            ).all()
        covers = {row.public_id: row.cover for row in snapshots}
        assert covers[first["items"][0]["revision_id"]] not in {None, "null", ""}
        assert covers[second["items"][0]["revision_id"]] in {None, "null"}
    assert second["items"][0]["revision_id"] != first["items"][0]["revision_id"]


@pytest.mark.asyncio
async def test_changed_entry_new_key_has_independent_identity_and_keeps_old_mapping(import_db):  # noqa: F811
    db = import_db
    external = uuid4().hex
    created, preview = await validated(db, source_package(external), "entry-first")
    first = await confirm(db, preview, "entry-save")
    old = await saved_entry(db, first["items"][0]["revision_id"])
    current = await mapped(db, external)
    changed, preview = await validated(
        db,
        source_package(
            external, meaning="银行", import_key="bank-finance", expected=current["current_version"]
        ),
        "entry-new-key",
    )
    assert await item_hash(db, changed["batch_id"]) != await item_hash(db, created["batch_id"])
    assert preview["items"][0]["operation"] == "UPDATE_DRAFT"
    revised = await confirm(db, preview, "entry-new-key-save")
    assert revised["items"][0]["status"] == "SUCCESS", revised
    new = await saved_entry(db, revised["items"][0]["revision_id"])
    assert new["chinese"] == "银行"
    assert all(new[k] != old[k] for k in ("entry_id", "sense_id", "source_id"))
    async with db.sessions() as session:
        rows = (
            (
                await session.execute(
                    text(
                        "SELECT import_entry_id,entry_id,sense_id,source_id "
                        "FROM import_entry_mapping "
                        "WHERE namespace='juya-content-library' AND content_type='dialogue' "
                        "AND content_id=:id ORDER BY import_entry_id"
                    ),
                    dict(id=first["items"][0]["content_id"]),
                )
            )
            .mappings()
            .all()
        )
    assert len(rows) == 2
    indexed = {r["import_entry_id"]: r for r in rows}
    assert all(indexed["bank-river"][k] == old[k] for k in ("entry_id", "sense_id", "source_id"))
    assert all(indexed["bank-finance"][k] == new[k] for k in ("entry_id", "sense_id", "source_id"))


@pytest.mark.asyncio
async def test_compat_docx_replay_never_merges_changed_input_by_title(import_db):  # noqa: F811
    db = import_db
    document = Document()
    document.add_heading("你好|Hello", level=1)
    document.add_paragraph("A: Hello.")
    document.add_paragraph("甲:你好。")

    def serialized():
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    async def create(key, data):
        result = await db.service.create(
            series_id=db.series,
            content_type="dialogue",
            mode="DOCX",
            data=data,
            actor=db.actor,
            key=key,
        )
        async with db.sessions() as session:
            stage = await session.scalar(
                text(
                    "SELECT public_id FROM content_import_stage "
                    "WHERE item_id=:id AND stage='VALIDATE'"
                ),
                dict(id=result["batch_id"]),
            )
        await db.service.validate(result["batch_id"], stage)
        return await db.service.preview(result["batch_id"])

    literal_bytes = serialized()
    preview = await create("compat-first", literal_bytes)
    first = await confirm(db, preview, "compat-save")
    assert first["items"][0]["status"] == "SUCCESS", first
    literal = await create("compat-literal", literal_bytes)
    assert literal["items"][0]["operation"] == "SKIP"
    document.paragraphs[1].text = "A: Good morning."
    changed = await create("compat-changed", serialized())
    assert changed["items"][0]["operation"] == "CREATE"
    second = await confirm(db, changed, "compat-new-save")
    assert second["items"][0]["status"] == "SUCCESS", second
    assert second["items"][0]["content_id"] != first["items"][0]["content_id"]
